"""Read only the official files and Q2/Q3 data interfaces; no Q4 imports."""
import sys
from pathlib import Path
sys.path.append(str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/Lib/site-packages'))
import openpyxl
import numpy as np
from pypdf import PdfReader
root = Path(__file__).resolve().parents[3]
for name in ['附件2.xlsx','附件3.xlsx','附件4.xlsx','result4-2.xlsx','result4-3.xlsx']:
    wb = openpyxl.load_workbook(root/'基础数据'/name, read_only=True, data_only=True)
    print(name)
    for ws in wb:
        print(ws.title, ws.max_row, ws.max_column)
        rows = list(ws.values)
        print([list(r[:8]) for r in rows[:8]])
        print('last:',list(rows[-1][:3]), list(rows[0][-3:]))
    wb.close()
for folder,names in [('q2_fresh_20260912',['fresh_forecasts.npz','main.npy']),('q3_revision_a_20260912',['conditional_forecasts.npz','solutions.npz'])]:
    for name in names:
        obj=np.load(root/'代码/实验'/folder/'results'/name)
        print(folder,name, {k:obj[k].shape for k in obj.files} if hasattr(obj,'files') else obj.shape)
for page in PdfReader(root/'基础数据/C题.pdf').pages:
    print(page.extract_text())
