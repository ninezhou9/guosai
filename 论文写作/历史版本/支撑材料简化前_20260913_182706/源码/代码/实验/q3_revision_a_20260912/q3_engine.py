"""Q3 route A: issue-time forecasts, paid plan revisions, stochastic DP.

Uses read-only Q2 source and its hash-verified causal load forecast artifact.
The published PV forecasts are treated as hourly point forecasts.
"""
from pathlib import Path
import sys
import importlib.util
import json
import hashlib
from datetime import datetime, timedelta
import numpy as np
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
Q2 = BASE.parent / 'q2_fresh_20260912'
try:
    import openpyxl
except ImportError:
    sys.path.append(str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/Lib/site-packages'))
    import openpyxl
spec = importlib.util.spec_from_file_location('q3_q2_core', Q2/'fresh_engine.py')
b = importlib.util.module_from_spec(spec)
spec.loader.exec_module(b)
DT, ETA, CAP, EMIN, EMAX, GRID = b.DT, b.ETA, b.CAP, b.EMIN, b.EMAX, b.GRID
QUANTILES = b.QUANTILES
TOL = 2e-5


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def load_inputs():
    data = ROOT/'基础数据'
    manifest = json.loads((Q2/'results/manifest.json').read_text(encoding='utf-8'))
    for name, digest in manifest['raw_inputs'].items():
        assert sha(data/name) == digest, f'Q2 input mismatch: {name}'
    for name, digest in manifest['source_hashes'].items():
        raw=(Q2/name).read_bytes()
        # Git on Windows converted LF to CRLF; accept only this exact byte change.
        normalized=hashlib.sha256(raw.replace(b'\r\n',b'\n')).hexdigest()
        assert sha(Q2/name) == digest or normalized == digest, f'Q2 source mismatch: {name}'
    cache = Q2/'results/fresh_forecasts.npz'
    assert sha(cache) == manifest['output_hashes']['fresh_forecasts.npz']
    with np.load(cache) as f:
        load_pred = f['separate'][:, 0].copy()
    wb = openpyxl.load_workbook(data/'附件2.xlsx', read_only=True, data_only=True)
    y = []
    for name in ['小区负载', '光伏发电实际功率']:
        rows = list(wb[name].values)
        assert len(rows) == 366 and len(rows[0]) == 145
        assert [r[0] for r in rows[1:]] == [datetime(2025,1,1)+timedelta(days=d) for d in range(365)]
        y.append(np.asarray([r[1:] for r in rows[1:]], float))
    wb.close()
    y = np.asarray(y)
    wb = openpyxl.load_workbook(data/'附件1.xlsx', read_only=True, data_only=True)
    p = np.array([r[1] for r in list(wb.active.values)[1:]], float)
    wb.close()
    forecasts = np.full((365,4,24), np.nan)
    wb = openpyxl.load_workbook(data/'附件3.xlsx', read_only=True, data_only=True)
    day = None
    for row in list(wb.active.values)[1:]:
        if row[0]: day = (datetime.strptime(row[0], '%Y-%m-%d')-datetime(2025,1,1)).days
        hour = int(row[1].split(':')[0]); assert hour in [0,6,12,18]
        issue = hour//6
        assert np.isnan(forecasts[day,issue]).all()
        forecasts[day,issue] = row[2:]
    wb.close()
    assert y.shape == (2,365,144) and p.shape == (144,)
    assert np.isfinite(y).all() and np.isfinite(forecasts).all() and (forecasts>=0).all()
    return y, p, load_pred, forecasts


def pv_interpolate(y, forecasts):
    pv = np.full((365,4,144), np.nan)
    for d in range(365):
        for k in range(4):
            start = k*36
            anchor = y[1,d,start-1] if start else (y[1,d-1,-1] if d else 0.)
            values = np.r_[anchor, forecasts[d,k]]
            pv[d,k,start:] = np.interp(np.arange(1,145-start)/6., np.arange(25), values)
    return pv


def conditional_forecasts(y, load_pred, forecasts):
    pv = pv_interpolate(y, forecasts)
    mu = load_pred[:,None,:]-pv
    net = y[0]-y[1]
    residual = net[:,None,:]-mu
    pred = np.full((365,4,5,144), np.nan)
    for d in range(1,365):
        ids = np.arange(max(14,d-30), d)
        for k in range(4):
            start = k*36
            pred[d,k,:,start:] = mu[d,k,start:][None,:]
            if len(ids)>=4:
                for lo in range(start,144,36):
                    hi = min(144,lo+36)
                    adjust = np.quantile(residual[ids,k,lo:hi], QUANTILES)
                    pred[d,k,:,lo:hi] += adjust[:,None]
            pred[d,k,:,start:] = np.sort(pred[d,k,:,start:],axis=0)
    return pred, pv


def model_day(y, pred, d, k, atoms=7):
    """Estimates use completed days only; initial class is a historical prior."""
    start = k*36; H = 144-start
    ids = np.arange(max(18,d-30),d)
    assert len(ids)>0
    hist = pred[ids,k,:,start:]
    scale_hist = np.maximum((hist[:,3]-hist[:,0])/1.683,100.)
    z = ((y[0]-y[1])[ids,start:]-hist[:,1])/scale_hist
    states = np.digitize(z,[-.84,.84])
    scale = np.maximum((pred[d,k,3,start:]-pred[d,k,0,start:])/1.683,100.)
    center = pred[d,k,1,start:]
    count = np.ones((3,3)); np.add.at(count,(states[:,:-1].ravel(),states[:,1:].ravel()),1)
    global_p = count/count.sum(axis=1,keepdims=True)
    P = np.empty((H-1,3,3)); samples = np.zeros((H,3,atoms)); weights = np.zeros_like(samples)
    for t in range(H):
        lo, hi = max(0,t-3), min(H,t+4)
        for s in range(3):
            values = z[:,lo:hi][states[:,lo:hi]==s]
            if not len(values): values = np.array([[-1.4,0.,1.4][s]])
            a,w = b.atomize(values,atoms)
            samples[t,s,:len(w)] = (center[t]+scale[t]*a)*DT
            weights[t,s,:len(w)] = w
        if t<H-1:
            lo, hi = max(0,t-3), min(H-1,t+4)
            count = 3*global_p.copy()
            np.add.at(count,(states[:,lo:hi].ravel(),states[:,lo+1:hi+1].ravel()),1)
            P[t] = count/count.sum(axis=1,keepdims=True)
    initial = np.bincount(states[:,0],minlength=3)+[.6,1.8,.6]
    initial = initial/initial.sum()
    return dict(P=P,samples=samples,weights=weights,initial=initial,center=center,scale=scale)


def revision_cost(new, old, p):
    diff = new-old
    up, down = np.maximum(diff,0), np.maximum(-diff,0)
    return float(1.5*p@up-.5*p@down)


def revision_lp(p, demand, E0, old):
    """No emergency in candidate LP. Includes refunds and cancellation charge."""
    H = len(p); dim = 7*H+1
    # g,c,b,w,E[H+1],up,down
    g,c,q,w,E,up,down = 0,H,2*H,3*H,4*H,5*H+1,6*H+1
    obj = np.zeros(dim); obj[up:up+H]=1.5*p;obj[down:down+H]=-.5*p
    bounds = [(0,None)]*dim
    A = lil_matrix((3*H,dim)); rhs = np.r_[demand,np.zeros(H),old]
    for t in range(H):
        A[t,[g+t,c+t,q+t,w+t]]=[1,-1,1,-1]
        A[H+t,[E+t+1,E+t,c+t,q+t]]=[1,-1,-ETA,1/ETA]
        A[2*H+t,[g+t,up+t,down+t]]=[1,-1,1]
        bounds[c+t]=bounds[q+t]=(0,CAP)
    for t in range(H+1):bounds[E+t]=(EMIN,EMAX)
    bounds[E]=(E0,E0)
    res=linprog(obj,A_eq=A.tocsr(),b_eq=rhs,bounds=bounds,method='highs')
    assert res.success,res.message
    assert abs(A@res.x-rhs).max()<TOL
    plan=np.maximum(res.x[:H],0)
    assert abs(revision_cost(plan,old,p)-res.fun)<TOL
    return plan


def interp_batch(f,x,grid):
    """f: K,G; x: K,S,G. Linear interpolation on a regular grid."""
    loc=np.clip((x-grid[0])/(grid[1]-grid[0]),0,len(grid)-1)
    i=np.minimum(loc.astype(int),len(grid)-2); frac=loc-i
    rows=np.arange(len(f))[:,None,None]
    return (1-frac)*f[rows,i]+frac*f[rows,i+1]


def bellman_batch(p,plans,m,grid=GRID):
    """Vectorized, identical piecewise-linear minimization to the Q2 core.

    Batch candidates share exogenous demand samples, not inventory or actions.
    """
    plans=np.atleast_2d(plans);K,H=plans.shape;G=len(grid)
    V=np.zeros((K,H+1,3,G)); step=grid[1]-grid[0]
    width=int(np.floor(CAP/ETA/step+1e-10))+1
    ix=np.arange(G)[:,None]-np.arange(width)[None,:]
    valid=ix>=0;ix=np.maximum(ix,0)
    for t in range(H-1,-1,-1):
        f_all=V[:,t+1] if t==H-1 else np.einsum('ij,kjg->kig',m['P'][t],V[:,t+1])
        for s in range(3):
            f=f_all[:,s];r=m['samples'][t,s][None,:]-plans[:,t,None]
            x=grid[None,None,:]
            distance=np.minimum(np.maximum(r,0),CAP)/ETA
            lo=np.maximum(EMIN,x-distance[:,:,None])
            knot=f+5*p[t]*ETA*grid
            windows=np.where(valid[None,:,:],knot[:,ix],np.inf)
            windows=np.minimum.accumulate(windows,axis=2)
            size=np.floor(distance/step+1e-10).astype(int)
            minima=windows[np.arange(K)[:,None,None],np.arange(G)[None,None,:],size[:,:,None]]
            boundary=interp_batch(f,lo,grid)+5*p[t]*ETA*lo
            deficit=5*p[t]*(r[:,:,None]-ETA*x)+np.minimum(minima,boundary)
            nxt=np.minimum(EMAX,x+ETA*np.minimum(np.maximum(-r,0),CAP)[:,:,None])
            surplus=interp_batch(f,nxt,grid)
            cost=np.where(r[:,:,None]>0,deficit,surplus)
            V[:,t,s]=np.einsum('ksg,s->kg',cost,m['weights'][t,s])
    assert np.isfinite(V).all() and np.diff(V,axis=3).max()<TOL
    return V


def day_run(y,p,pred,models,d,E0,mask=7,update_only=False,grid=GRID,return_scores=False):
    initial=b.plan_lp(p,pred[d,0,3]*DT,E0,0.)
    current=initial.copy();E=float(E0);rows=[];revisions=[];scores=[]
    active=0;start=0;model=models[0]
    V=bellman_batch(p,current,model,grid)[0]
    for t in range(144):
        if t in (36,72,108):
            k=t//36
            if update_only or (mask & (1<<(k-1))):
                active=k;start=t;model=models[k];old=current[t:].copy()
                plans=[old];labels=['keep']
                if not update_only:
                    for qi in [1,2,3,4]:
                        proposal=revision_lp(p[t:],pred[d,k,qi,t:]*DT,E,old)
                        if not any(np.allclose(proposal,g,atol=1e-8,rtol=0) for g in plans):
                            plans.append(proposal);labels.append(str(QUANTILES[qi]))
                plans=np.asarray(plans);Vs=bellman_batch(p[t:],plans,model,grid)
                charge=np.array([revision_cost(g,old,p[t:]) for g in plans])
                expected=np.array([sum(model['initial'][s]*np.interp(E,grid,v[0,s]) for s in range(3)) for v in Vs])
                total=charge+expected;best=int(np.argmin(total))
                assert total[best]<=total[0]+TOL
                current[t:]=plans[best];V=Vs[best]
                diff=current[t:]-old
                for j,(before,after,delta) in enumerate(zip(old,current[t:],diff),start=t):
                    if abs(delta)>1e-8:
                        up=max(delta,0.);down=max(-delta,0.)
                        revisions.append([k,j,before,after,up,down,1.5*p[j]*up,p[j]*down,.5*p[j]*down])
                if return_scores:
                    for l,c,e,v in zip(labels,charge,expected,total):scores.append([k,l,float(c),float(e),float(v),l==labels[best]])
        loc=t-start
        actual=(y[0,d,t]-y[1,d,t]);s=int(np.digitize((actual-model['center'][loc])/model['scale'][loc],[-.84,.84]))
        future=V[loc+1,s] if t==143 else model['P'][loc,s]@V[loc+1]
        c,q,e,w,nxt=b.action(E,actual*DT-current[t],p[t],future,grid)
        rows.append([current[t],c,q,e,w,E,nxt]);E=nxt
    arr=np.asarray(rows);b.validate(arr,(y[0]-y[1])[d],p)
    return dict(arr=arr,initial=initial,ledger=np.asarray(revisions,float).reshape(-1,9),scores=scores)


def summary_day(run,p):
    arr=run['arr'];ledger=run['ledger']
    initial=float(run['initial']@p)
    up=float(ledger[:,6].sum());refund=float(ledger[:,7].sum());penalty=float(ledger[:,8].sum())
    emergency=float(5*p@arr[:,3])
    return dict(initial_cost=initial,increase_cost=up,cancellation_refund=refund,cancellation_penalty=penalty,
                adjustment_net=up-refund+penalty,emergency_cost=emergency,total_cost=initial+up-refund+penalty+emergency,
                initial_kwh=float(run['initial'].sum()),effective_kwh=float(arr[:,0].sum()),
                increase_kwh=float(ledger[:,4].sum()),cancel_kwh=float(ledger[:,5].sum()),
                emergency_kwh=float(arr[:,3].sum()),start_kwh=float(arr[0,5]),end_kwh=float(arr[-1,6]))


def warmup(y,p):
    E=6000.;rows=[]
    net=y[0]-y[1]
    for d in range(31):
        plan=np.zeros(144) if d==0 else np.maximum(net[d-1]*DT,0)
        a=b.execute(net[d],p,plan,E,None,None,greedy=True);rows.append(a);E=float(a[-1,6])
    return np.asarray(rows)


def date(d):
    return (datetime(2025,1,1)+timedelta(days=int(d))).strftime('%Y-%m-%d')
