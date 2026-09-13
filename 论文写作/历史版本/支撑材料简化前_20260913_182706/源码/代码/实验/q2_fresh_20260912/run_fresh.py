"""原始附件起步：每次完整重训，不加载任何预测或调度缓存。"""
from pathlib import Path
from datetime import datetime,timedelta
import argparse,csv,hashlib,json,sys,time,platform
import numpy as np
import openpyxl
import fresh_engine as m

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[2]

def dump(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def log(*x):print(datetime.now().isoformat(timespec='seconds'),*x,flush=True)
def date(d):return (datetime(2025,1,1)+timedelta(days=int(d))).strftime('%Y-%m-%d')
def csvout(path,heads,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(heads);w.writerows(rows)

def load(data):
    wb=openpyxl.load_workbook(data/'附件2.xlsx',read_only=True,data_only=True)
    arrays=[];headers=[]
    for name in ['小区负载','光伏发电实际功率']:
        rows=list(wb[name].values)
        assert len(rows)==366 and all(len(row)==145 for row in rows)
        assert [row[0] for row in rows[1:]]==[datetime(2025,1,1)+timedelta(days=d) for d in range(365)]
        arrays.append(np.asarray([row[1:] for row in rows[1:]],float));headers.append(rows[0][1:])
    wb.close()
    wb=openpyxl.load_workbook(data/'附件1.xlsx',read_only=True,data_only=True)
    rows=list(wb.active.values);wb.close()
    def minute(x):
        if hasattr(x,'hour'):return x.hour*60+x.minute
        if x=='0:00+1':return 1440
        h,v=map(int,str(x).split(':'));return h*60+v
    assert len(rows)==145
    assert [minute(x) for x in headers[0]]==[minute(x) for x in headers[1]]==[minute(r[0]) for r in rows[1:]]==list(range(10,1441,10))
    y=np.asarray(arrays);p=np.asarray([r[1] for r in rows[1:]],float)
    assert y.shape==(2,365,144) and np.isfinite(y).all() and y.min()>=0
    assert p.shape==(144,) and np.isfinite(p).all() and p.min()>0
    return y,p

def forecast_fresh(y,out):
    X=np.zeros((365,144,25))
    for d in range(1,365):X[d]=m.features(y,d)
    net=y[0]-y[1];direct=np.zeros((365,5,144));separate=np.zeros((365,2,144))
    base=np.zeros_like(direct);pred=np.zeros_like(direct);family='direct';selection={};trace=[]
    for d in range(1,365):
        if d<14:
            direct[d]=net[d-1][None,:];separate[d]=y[:,d-1]
        else:
            # 必须调用真实拟合，没有exists/cache/read预测文件分支。
            direct[d],separate[d]=m.forecast_day(y,d,X)
            trace.append([date(d),date(max(7,d-56)),date(d-1),7,m.FIT_CALLS])
        if d==21:
            a=float(abs(direct[14:21,1]-net[14:21]).mean())
            b=float(abs(separate[14:21,0]-separate[14:21,1]-net[14:21]).mean())
            family='direct' if a<=b else 'separate_center'
            selection=dict(validation='2025-01-15..2025-01-21',frozen_at='2025-01-22 00:00',
                           direct_MAE_kW=a,separate_MAE_kW=b,family=family)
            log('forecast family',selection)
        base[d]=direct[d]
        if family=='separate_center':base[d]+=separate[d,0]-separate[d,1]-direct[d,1]
        pred[d]=base[d]
        ids=np.arange(max(14,d-30),d)
        if len(ids)>=4:
            for block in range(4):
                sl=slice(block*36,(block+1)*36)
                for k,q in enumerate(m.QUANTILES):
                    pred[d,k,sl]+=np.quantile(net[ids,sl]-base[ids,k,sl],float(q))
            pred[d]=np.sort(pred[d],axis=0)
        if d%15==0 or d==364:
            log('fresh forecast',date(d),'fitted_models',m.FIT_CALLS)
    assert m.FIT_CALLS==351*7
    np.savez_compressed(out/'fresh_forecasts.npz',pred=pred,base=base,direct=direct,separate=separate)
    csvout(out/'training_trace.csv',['forecast_date','training_start','training_end','models_fit','cumulative_models'],trace)
    dump(out/'forecast_selection.json',selection)
    return pred,selection

def warmup(y,p):
    net=y[0]-y[1];E=6000.;days=[]
    for d in range(31):
        g=np.zeros(144) if d==0 else np.maximum(net[d-1]/6,0)
        a=m.execute(net[d],p,g,E,None,None,greedy=True);days.append(a);E=float(a[-1,-1])
    return np.asarray(days)

def select_kappa(y,p,pred,warm,out):
    ref=float(.9*p.mean());grid=ref*np.array([0,.25,.5,.75,1,1.25]);rows=[]
    for k in grid:
        E=float(warm[20,-1,-1]);start=E;days=[]
        for d in range(21,31):
            model=m.markov(y,pred,d)
            g=m.plan_lp(p,pred[d,3]/6,E,0.)
            V=m.bellman(p,g,model,k)
            a=m.execute((y[0]-y[1])[d],p,g,E,model,V);days.append(a);E=float(a[-1,-1])
        a=np.asarray(days);cost=float(np.sum(a[:,:,0]*p+5*a[:,:,3]*p))
        rows.append(dict(kappa=float(k),bill=cost,inventory_adjusted=cost+ref*(start-E),start=start,end=E))
    best=min(rows,key=lambda x:(x['inventory_adjusted'],x['kappa']))
    record=dict(candidates=rows,selected=best,reference_inventory_price=ref,
                validation='2025-01-22..2025-01-31',frozen_at='2025-02-01 00:00',
                model='fixed q80, day-ahead LP terminal coefficient zero, own continuous SOC')
    dump(out/'controller_selection.json',record);log('new controller selection',record)
    return best['kappa'],record

def backtest(y,p,pred,warm,kappa):
    E0=float(warm[-1,-1,-1]);states={name:E0 for name in ['main','same_plan_greedy','independent_greedy']}
    runs={name:[] for name in states};net=y[0]-y[1]
    for d in range(31,365):
        model=m.markov(y,pred,d);k=0. if d==364 else kappa
        g=m.plan_lp(p,pred[d,3]/6,states['main'],0.)
        V=m.bellman(p,g,model,k)
        main=m.execute(net[d],p,g,states['main'],model,V)
        pair=m.execute(net[d],p,g,states['same_plan_greedy'],None,None,greedy=True)
        gi=m.plan_lp(p,pred[d,3]/6,states['independent_greedy'],0.)
        independent=m.execute(net[d],p,gi,states['independent_greedy'],None,None,greedy=True)
        for name,a in zip(states,[main,pair,independent]):runs[name].append(a);states[name]=float(a[-1,-1])
        if (d-31)%50==0 or d==364:log('fresh backtest',date(d))
    return {name:np.asarray(a) for name,a in runs.items()}

def verify(y,p,pred,runs,kappa):
    checks={name:m.validate(a.reshape(-1,7),(y[0]-y[1])[31:].ravel(),np.tile(p,334)) for name,a in runs.items()}
    assert np.array_equal(runs['main'][:,:,0],runs['same_plan_greedy'][:,:,0])
    changed=y.copy();changed[:,31:]+=10000
    assert np.array_equal(m.features(y,31),m.features(changed,31))
    model=m.markov(y,pred,31);other=m.markov(changed,pred,31)
    for key in model:np.testing.assert_array_equal(model[key],other[key])
    g=runs['main'][0,:,0];V=m.bellman(p,g,model,kappa);actual=(y[0]-y[1])[31]
    modified=actual.copy();modified[72:]+=20000
    a=m.execute(actual,p,g,runs['main'][0,0,5],model,V)
    b=m.execute(modified,p,g,runs['main'][0,0,5],model,V)
    np.testing.assert_array_equal(a[:72],b[:72])
    # 类内样本均值不能替代真实缺口函数。
    toy=dict(samples=np.tile([0.,100.],(1,3,1)),weights=np.full((1,3,2),.5),P=np.empty((0,3,3)))
    values=m.bellman(np.array([1.]),np.array([50.]),toy,0.)
    assert abs(values[0,0,0]-125)<1e-8
    return dict(physical=checks,dayahead_causal_features_and_probabilities=True,
                future_observation_perturbation=True,within_class_shortfall_counterexample=125.)

def export(y,p,pred,runs,out,summary,kappa):
    heads=['date','slot','g_kwh','charge_kwh','discharge_kwh','emergency_kwh','waste_kwh','start_kwh','end_kwh','price','net_kwh']
    all_daily=[];dates=[date(d) for d in range(31,365)]
    for name,a in runs.items():
        np.save(out/f'{name}.npy',a)
        rows=((dates[i],t+1,*a[i,t],float(p[t]),float((y[0,i+31,t]-y[1,i+31,t])/6)) for i in range(334) for t in range(144))
        path=out/f'{name}_intervals.csv';csvout(path,heads,rows)
        with path.open(encoding='utf-8-sig',newline='') as f:read=list(csv.DictReader(f))
        nc=sum(float(r['g_kwh'])*float(r['price']) for r in read)
        ec=sum(5*float(r['emergency_kwh'])*float(r['price']) for r in read)
        assert len(read)==48096 and abs(nc-summary[name]['normal_cost'])<1e-5 and abs(ec-summary[name]['emergency_cost'])<1e-5
        for i,ds in enumerate(dates):
            normal=float(a[i,:,0]@p);em=float(5*a[i,:,3]@p)
            all_daily.append([name,ds,normal,em,normal+em,float(a[i,:,3].sum()),float(a[i,0,5]),float(a[i,-1,-1])])
    csvout(out/'daily_summary.csv',['policy','date','normal_cost','emergency_cost','total_cost','emergency_kwh','start_kwh','end_kwh'],all_daily)
    wb=openpyxl.Workbook(write_only=True);main=runs['main']
    ws=wb.create_sheet('计划购电量');ws.append(['日期']+[f'{t*10//60:02d}:{t*10%60:02d}-{(t+1)*10//60:02d}:{(t+1)*10%60:02d}' for t in range(144)]+['全天购电量','全天购电费'])
    for i,ds in enumerate(dates):ws.append([ds,*main[i,:,0].tolist(),float(main[i,:,0].sum()),float(main[i,:,0]@p)])
    ws=wb.create_sheet('十分钟明细');ws.append(heads)
    for i,ds in enumerate(dates):
        for t in range(144):ws.append([ds,t+1,*main[i,t].tolist(),float(p[t]),float((y[0,i+31,t]-y[1,i+31,t])/6)])
    ws=wb.create_sheet('四个指定日');ws.append(heads)
    for ds in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
        i=dates.index(ds)
        for t in range(144):ws.append([ds,t+1,*main[i,t].tolist(),float(p[t]),float((y[0,i+31,t]-y[1,i+31,t])/6)])
    ws=wb.create_sheet('充放电量');ws.append(['日期','四小时时段','充电量kWh','放电量kWh','日初库存kWh','日末库存kWh'])
    for i,ds in enumerate(dates):
        for block in range(6):ws.append([ds,f'{block*4:02d}:00-{(block+1)*4:02d}:00',float(main[i,block*24:(block+1)*24,1].sum()),float(main[i,block*24:(block+1)*24,2].sum()),float(main[i,0,5]),float(main[i,-1,-1])])
    ws=wb.create_sheet('紧急购电量');ws.append(['日期','时段序号','应急电量kWh','应急费用元'])
    for i,ds in enumerate(dates):
        for t in range(144):
            if main[i,t,3]>1e-8:ws.append([ds,t+1,float(main[i,t,3]),float(5*p[t]*main[i,t,3])])
    ws=wb.create_sheet('说明');ws.append(['项目','内容'])
    for key,value in dict(forecast='原始附件完整重训2457个LightGBM模型；未读取旧预测',quantile=.8,lp_terminal=0,dp_terminal=kappa,dec31_terminal=0,observation='当前时段功率即时可测',period='2025-02-01..2025-12-31',status='独立审核结果，区间终点按00:10至24:00解释').items():ws.append([key,str(value)])
    wb.save(out/'result2.xlsx')
    wb=openpyxl.load_workbook(out/'result2.xlsx',read_only=True,data_only=True)
    rr=list(wb['计划购电量'].values);assert len(rr)==335
    assert abs(sum(float(row[-1]) for row in rr[1:])-summary['main']['normal_cost'])<1e-5
    rr=list(wb['十分钟明细'].values);assert len(rr)==48097
    assert abs(sum(5*float(row[5])*float(row[9]) for row in rr[1:])-summary['main']['emergency_cost'])<1e-5
    wb.close()

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data-dir',type=Path)
    parser.add_argument('--output',type=Path,default=BASE/'results')
    args=parser.parse_args();out=args.output.resolve()
    if out.exists() and any(out.iterdir()):raise RuntimeError('输出目录非空；请选择新的--output，禁止复用缓存或覆盖旧结果')
    out.mkdir(parents=True,exist_ok=True)
    data=args.data_dir or next(p for p in [ROOT/'基础数据/C题/附件',ROOT/'基础数据'] if (p/'附件2.xlsx').exists())
    started=time.perf_counter();y,p=load(data)
    input_hashes={file:sha(data/file) for file in ['附件1.xlsx','附件2.xlsx']}
    log('raw inputs read; no prior prediction or results loaded',input_hashes)
    pred,forecast_selection=forecast_fresh(y,out)
    warm=warmup(y,p);np.save(out/'january_warmup.npy',warm)
    kappa,controller_selection=select_kappa(y,p,pred,warm,out)
    runs=backtest(y,p,pred,warm,kappa)
    checks=verify(y,p,pred,runs,kappa)
    summary={name:m.metrics(a,p) for name,a in runs.items()}
    err=(y[0]-y[1])[31:]-pred[31:,1]
    forecast_metrics=dict(MAE_kW=float(abs(err).mean()),RMSE_kW=float(np.sqrt((err**2).mean())),q80_coverage=float(np.mean((y[0]-y[1])[31:]<=pred[31:,3])))
    export(y,p,pred,runs,out,summary,kappa)
    manifest=dict(created=datetime.now().isoformat(),status='completed_checked_pending_review',raw_inputs=input_hashes,
                  old_prediction_reads=0,old_result_reads=0,models_trained=m.FIT_CALLS,
                  elapsed_seconds=time.perf_counter()-started,summary=summary,checks=checks,
                  forecast_metrics=forecast_metrics,forecast_selection=forecast_selection,controller_selection=controller_selection,
                  python=sys.version,numpy=np.__version__,scipy=m.scipy.__version__,lightgbm=m.lgb.__version__,openpyxl=openpyxl.__version__,
                  source_hashes={f:sha(BASE/f) for f in ['fresh_engine.py','run_fresh.py']},seed=m.SEED,
                  command=[sys.executable,str(Path(__file__).relative_to(ROOT))],
                  output_hashes={str(f.relative_to(out)):sha(f) for f in out.iterdir() if f.is_file()})
    dump(out/'manifest.json',manifest)
    dump(out/'summary.json',dict(summary=summary,forecast_metrics=forecast_metrics,kappa=kappa,checks=checks))
    log('COMPLETE',json.dumps(summary,ensure_ascii=False),'seconds',manifest['elapsed_seconds'])

if __name__=='__main__':main()
