"""去除既有流程图公式，并按解题信息结构生成第三、四问可编辑 PPT。

依赖：python-pptx、lxml。导出 PNG/PDF 使用同目录的导出核验.py。
第一、二问直接修改 OOXML 文本，保留原有框、箭头、几何与文字样式。
第三、四问采用原参考图的颜色、线宽、箭头和节点尺寸；仅扩展画布、重排分组。
"""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def replace_text(shape, text):
    """同时替换 Office 数学主分支及兼容分支，避免隐藏公式残留。"""
    body = shape.find("p:txBody", NS)
    if body is None:
        return
    old = body.findall("a:p", NS)
    for p in old:
        body.remove(p)
    for i, line in enumerate(text.split("\n")):
        source = old[min(i, len(old) - 1)] if old else None
        p = ET.SubElement(body, "{%s}p" % NS["a"])
        prop = source.find("a:pPr", NS) if source is not None else None
        if prop is not None:
            p.append(deepcopy(prop))
        r = ET.SubElement(p, "{%s}r" % NS["a"])
        rp = source.find("a:r/a:rPr", NS) if source is not None else None
        if rp is None and prop is not None:
            rp = prop.find("a:defRPr", NS)
        if rp is not None:
            rp = deepcopy(rp)
            rp.tag = "{%s}rPr" % NS["a"]
            # 原数学对象改用与正文一致的中文字体；字号、颜色不变。
            for font in rp:
                if font.get("typeface") == "Cambria Math":
                    font.set("typeface", "SimSun")
            r.append(rp)
        ET.SubElement(r, "{%s}t" % NS["a"]).text = line
    for font in body.iter():
        if font.get("typeface") == "Cambria Math":
            font.set("typeface", "SimSun")


Q1_REFERENCE = {
    "net-load": "净负荷\n负荷扣光伏",
    "time-grid": "时间离散\n十分钟一段",
    "energy-conversion": "功率转电量\n按时段长度换算为千瓦时",
    "milp-objective-vars": "购电费用最小\n联合确定\n购电与充放电\n储电量与弃光\n充放电互斥状态",
    "cyclic-milp": "构建日循环混合整数规划\n日初与日末库存相同\n共同库存由优化确定\n允许弃光，不售电",
    "bottom-note": "已知预测条件下优化日循环储能；初末库存相同且共同优化；不设紧急购电与售电",
}
Q1_REBUILD = {
    "eq-known-vars": "已知电价与供需预测",
    "eq-timestep": "每段十分钟",
    "eq-netload": "负荷扣除光伏后换算电量",
    "eq-objective": "最小化全天购电费用",
    "eq-balance": "电力供给满足负荷、充电与弃光",
    "eq-soc": "按充放电效率逐段更新库存",
    "eq-bounds": "储电量与充放电功率均受上下限约束",
    "eq-cyclic": "初末库存相同，充电与放电互斥",
    "baseline-label": "基线购电规则",
    "eq-baseline": "净需求缺口由电网补足",
    "optimal-schedule": "最优调度\n购电与储能\n弃光轨迹",
    "baseline": "无储能\n购电基线\n松弛下界",
    "physical-checks": "逐时核验\n能量守恒\n库存与功率",
}
Q1_LOGICAL = {
    "eq-input": "已知电价与供需预测",
    "eq-time": "全天等间隔划分",
    "eq-netload": "负荷扣除光伏\n按时段换算电量",
    "eq-objective": "联合确定各时段购电与储能安排",
    "eq-balance": "电力供给满足负荷、充电与弃光",
    "eq-soc": "按充放电效率逐段更新储电量",
    "eq-mutual": "充电与放电不同时发生，功率受限",
    "eq-bound": "储电量始终处于允许范围",
    "eq-boundary": "日初与日末库存相同，充放电状态互斥",
    "eq-baseline": "净需求缺口由电网补足",
}
Q2_METHOD = {
    "10": "全天购电计划固定不变",
    "30": "以第八十分位净需求为输入，最小化购电费用；满足库存与功率约束，输出全天普通购电计划。",
    "35": "标准化误差按上下阈值分为三态；全局计数平滑，再用相邻时段局部计数修正；类内最多七个加权样本。",
    "40": "在库存网格上计算价值函数，终端价值设为零；每十分钟读取实测，富余时充电或弃用，缺口时择优放电并紧急购电。",
    "43": "储电量与充放电功率受边界约束，计入充放电损耗；紧急购电按五倍电价结算。决策只用当前与过去观测，日末库存连续传递至次日。",
    "46": "全天普通购电计划、充放电轨迹、紧急购电记录、逐时库存与逐日费用，用于全年账单核算与方案对比。",
}
Q2_ROUTE = {
    "37": "日前线性规划\n最小化购电费用\n满足储电量边界\n限制充放电功率",
    "39": "全天购电计划\n日内不再修改",
    "41": "三状态误差模型\n按上下阈值划分类别",
    "43": "估计状态转移\n全局平滑与局部修正",
    "47": "终端价值参数验证\n依据一月验证期择定",
    "51": "日内反馈执行\n每十分钟读取实测\n富余时优先充电\n缺口时放电并紧急购电",
    "53": "更新储电量\n计入充放电损耗\n日末库存传递至次日",
    "55": "逐日费用核算\n普通购电按分时电价\n紧急购电按五倍电价",
}


def clean_existing():
    for path in sorted(ROOT.rglob("*.pptx")):
        if "问题一" in path.parts:
            replacements = (Q1_LOGICAL if "逻辑重排" in path.name else
                            Q1_REBUILD if "重构" in path.name else Q1_REFERENCE)
            by_id = False
        elif "问题二" in path.parts and "参考风格" in path.name:
            replacements = {"bellman-policy": "Bellman 动态规划\n计算固定计划下库存价值\n最小化预期紧急购电费\n终端价值设为零"}
            by_id = False
        elif "问题二" in path.parts and "旧版备份" in path.name:
            replacements = {
                "bellman-value-body": "计算库存价值函数\n终端价值设为零",
                "accounting-verification-body": "核算计划购电费与紧急购电费\n检查功率平衡、储能边界与库存连续性\n输出购电计划、充放电轨迹与紧急购电记录",
                "feedback-action-body": "富余时充电，余量弃用\n缺口时放电并紧急购电\n按价值函数选择放电量",
            }
            by_id = False
        elif "问题二" in path.parts and "方法流程图" in path.name:
            replacements, by_id = Q2_METHOD, True
        elif "问题二" in path.parts and path.name == "第二问_技术路线图.pptx":
            replacements, by_id = Q2_ROUTE, True
        else:
            continue
        with ZipFile(path) as z:
            files = {name: z.read(name) for name in z.namelist()}
        for name in files:
            if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                continue
            tree = ET.fromstring(files[name])
            for shape in tree.findall(".//p:sp", NS):
                nv = shape.find("p:nvSpPr/p:cNvPr", NS)
                key = nv.get("id" if by_id else "name")
                if key in replacements:
                    replace_text(shape, replacements[key])
            files[name] = ET.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(path, "w", ZIP_DEFLATED) as z:
            for name, data in files.items():
                z.writestr(name, data)
        if "问题二" in path.parts and by_id:
            # 原 .slide 导出的包可被 python-pptx 读取，但部分 Office
            # 导入器无法识别；重新封装保留所有形状几何及 SVG 箭头层。
            Presentation(path).save(path)
        print("文字化：", path.relative_to(ROOT))


# 与原第三问参考 PPT 的 DrawingML 相同（设计像素为 96 dpi）。
STYLE = {"ink": "111111", "heading": "367CB3", "arrow": "8198B3",
         "gray": "777777", "group": "B0B0B0", "dash": "355C96",
         "input": "9EC3E6", "output": "BBD6EE"}
PX = 9525


class Diagram:
    def __init__(self, title, height):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = 1600 * PX, height * PX
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.group("overall", 20, 20, 1560, height - 65, outer=True)
        self.text("title", title, 40, 28, 1520, 41, 18.75, "heading", True)

    def text(self, name, text, x, y, w, h, size=15, color="ink", bold=False, shape=None):
        s = shape or self.slide.shapes.add_textbox(x*PX, y*PX, w*PX, h*PX)
        s.name = name
        tf = s.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = 9*PX if shape else 0
        tf.margin_top = tf.margin_bottom = 6*PX if shape else 0
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            p.space_before = p.space_after = Pt(0)
            r = p.add_run()
            r.text = line
            r.font.name, r.font.size = "SimSun", Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(STYLE[color])
            for tag in ("a:ea", "a:cs"):
                font = OxmlElement(tag)
                font.set("typeface", "SimSun")
                r._r.get_or_add_rPr().append(font)
        return s

    def node(self, name, text, x, y, kind="wide", fill=None):
        # 沿用原图四种节点尺寸，文字多少不影响框尺寸。
        w, h, size = {"wide": (320, 76, 15), "tall": (320, 87, 14.25),
                      "small": (169, 89, 15), "small2": (166, 89, 14.25)}[kind]
        s = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x*PX, y*PX, w*PX, h*PX)
        for style in s._element.findall("p:style", NS):
            s._element.remove(style)
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(STYLE[fill] if fill else "FFFFFF")
        s.line.color.rgb = RGBColor.from_string(STYLE[fill] if fill else STYLE["ink"])
        s.line.width = 18098
        s._element.spPr.append(OxmlElement("a:effectLst"))
        self.text(name, text, x, y, w, h, size, shape=s)
        return s

    def group(self, name, x, y, w, h, title=None, outer=False):
        s = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x*PX, y*PX, w*PX, h*PX)
        for style in s._element.findall("p:style", NS):
            s._element.remove(style)
        s.name = name
        s.fill.background()
        s.line.color.rgb = RGBColor.from_string(STYLE["arrow" if outer else "group"])
        s.line.width = 28575 if outer else 23813
        s._element.spPr.append(OxmlElement("a:effectLst"))
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        s._element.spPr.find("a:ln", NS).append(dash)
        if title:
            self.text(name+"-heading", title, x+10, y+9, w-20, 41, 18.75, "heading", True)

    def arrow(self, name, points, color="arrow"):
        for i, ((x1,y1),(x2,y2)) in enumerate(zip(points, points[1:])):
            s = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1*PX,y1*PX,x2*PX,y2*PX)
            for style in s._element.findall("p:style", NS):
                s._element.remove(style)
            s.name = name+f"-{i}"
            s.line.width = 20003
            s._element.spPr.append(OxmlElement("a:effectLst"))
            s.line.color.rgb = RGBColor.from_string(STYLE[color])
            if i == len(points)-2:
                end = OxmlElement("a:tailEnd")
                end.set("type", "triangle")
                end.set("w", "med")
                end.set("len", "med")
                s._element.spPr.find("a:ln", NS).append(end)

    def save(self, name):
        path = ROOT / name
        self.prs.save(path)
        print("重绘：", name)


def draw_q3():
    d = Diagram("第三问  多时刻预报驱动的购电调整与滚动反馈", 920)
    d.group("dayahead",40,80,1520,185,"零点：生成初始计划与执行所需的库存价值")
    for name,text,x in [
        ("initial-information","历史负荷与零点光伏预报\n只使用当时可获得的信息",60),
        ("conditional-forecast","条件净负荷分位预测\n光伏插值与历史净误差校准",440),
        ("initial-lp","日前线性规划\n八十分位需求与分时电价\n输入本策略日初库存",820),
        ("initial-policy","初始有效购电计划\n三状态动态规划计算库存价值",1200)]:
        d.node(name,text,x,145)
    for x in (380,760,1140):
        d.arrow("dayahead-link",[(x,183),(x+60,183)])
    d.group("revision",40,300,1120,250,"日内发布时刻：只调整尚未执行的购电计划")
    d.node("published-update","六、十二、十八时\n读取最新预报与当前库存\n更新条件预测和三状态误差模型",60,375,"tall")
    d.node("candidate-lp","多分位线性规划生成调整候选\n冻结已执行时段\n保留原计划作为不调整候选",440,375,"tall")
    d.node("candidate-dp","动态规划比较候选\n综合调整费与预期紧急购电费\n选定有效计划及对应库存价值",820,375,"tall")
    d.arrow("update-to-candidates",[(380,418),(440,418)])
    d.arrow("candidates-to-evaluation",[(760,418),(820,418)])
    d.text("keep-note","保持原计划时，仍使用更新后的预测与库存价值",70,485,1080,36,15,"gray")
    d.group("execution",1180,300,380,410,"每十分钟：读取实测并执行")
    d.node("current-observation","当前净负荷与实际库存\n结合有效计划和库存价值选择动作",1200,355)
    d.node("surplus","富余时段\n优先充电\n余量弃用",1200,485,"small")
    d.node("deficit","缺口时段\n择优放电\n紧急购电补足",1390,485,"small2")
    d.node("inventory","按实际充放电更新库存\n逐时记录，日末库存传递至次日",1200,610)
    d.arrow("initial-to-execute",[(1360,221),(1360,280),(1545,280),(1545,393),(1520,393)])
    d.arrow("revision-to-execute",[(1140,418),(1168,418),(1168,393),(1200,393)])
    d.arrow("surplus-branch",[(1360,431),(1360,455),(1284,455),(1284,485)],"gray")
    d.arrow("deficit-branch",[(1360,431),(1360,455),(1473,455),(1473,485)],"gray")
    d.arrow("surplus-to-stock",[(1284,574),(1284,591),(1360,591),(1360,610)],"gray")
    d.arrow("deficit-to-stock",[(1473,574),(1473,591),(1360,591)],"gray")
    d.arrow("next-issue-feedback",[(1200,648),(1170,648),(1170,580),(220,580),(220,462)])
    d.text("feedback-label","到启用的预报时刻，携带实际库存进入下一次调整",330,587,760,28,15,"heading")
    d.group("verification",40,630,1120,200,"结算与验证：保留每次调整的完整账本")
    d.node("ledger","分别记录初始购电与紧急购电\n逐笔记录增购、撤销退款及违约费",60,705)
    d.node("checks","独立复算费用与物理约束\n核对计划重放和库存连续\n结果表保存后回读",440,705)
    d.node("results","第三问结果表与逐时轨迹\n八种更新组合及固定计划对照",820,705,fill="output")
    d.arrow("ledger-to-checks",[(380,743),(440,743)])
    d.arrow("checks-to-output",[(760,743),(820,743)])
    d.arrow("records-to-ledger",[(1360,686),(1360,850),(220,850),(220,781)])
    d.text("ledger-label","逐时轨迹与逐笔交易记录",560,789,530,28,15,"heading")
    d.text("footer","分时电价已知；更新预报、计划调整与十分钟执行按时间先后衔接；跨日库存连续",40,881,1520,28,15,"gray")
    d.save("问题三/第三问_技术路线图_参考风格.pptx")


def draw_q4():
    d = Diagram("第四问  随机电价下的联合预测、双分支调度与真实结算",1040)
    d.group("joint-input",40,80,1520,205,"共同输入与联合模型：预测价用于规划，历史配对样本用于动态规划")
    items = [
        ("source","已核验的第二、三问预测\n历史负荷、光伏与真实电价",60),
        ("price-forecast","因果电价预测\n仅使用已完成时段的价格\n在允许的发布时间更新",440),
        ("paired-errors","净负荷与价格误差配对\n保留同一历史样本的对应关系\n构造带权联合情景",820),
        ("joint-model","九状态联合模型\n净负荷与价格类别共同描述状态\n估计转移概率与类内样本",1200)]
    for name,text,x in items:
        d.node(name,text,x,150,"tall")
    for x in (380,760,1140):
        d.arrow("joint-link",[(x,193),(x+60,193)])
    d.group("fixed-branch",40,335,740,355,"4-2：继承第二问，普通购电计划全天固定")
    d.node("fixed-input","零点净负荷分位预测\n零点预测价与本策略日初库存",60,405)
    d.node("fixed-plan","零点线性规划\n确定全天普通购电计划并冻结",440,405)
    d.node("fixed-dp","九状态联合动态规划\n在固定计划下计算库存价值\n综合未来需求与电价的不确定性",60,550,"tall")
    d.node("fixed-execution","每十分钟反馈执行\n读取当前净负荷、真实价与库存\n富余充电；缺口择优放电并补购",440,550,"tall")
    d.arrow("fixed-input-plan",[(380,443),(440,443)])
    d.arrow("fixed-plan-dp",[(600,481),(600,515),(220,515),(220,550)])
    d.arrow("fixed-dp-execution",[(380,593),(440,593)])
    d.group("adjustable-branch",800,335,760,355,"4-3：继承第三问，日内预报更新并允许调整")
    d.node("adjustable-initial","零点条件净负荷与预测价\n线性规划制定初始购电计划",820,405)
    d.node("adjustable-update","六、十二、十八时\n更新光伏预报与价格预测\n重新构建联合误差模型",1200,405)
    d.node("adjustable-execution","选定剩余时段计划及库存价值\n每十分钟按当前实测反馈执行\n更新库存并逐笔记录调整",820,550,"tall")
    d.node("adjustable-candidates","生成调整候选并保留原计划\n联合动态规划评价各候选\n比较调整费与预期紧急购电费",1200,550,"tall")
    d.arrow("initial-update",[(1140,443),(1200,443)])
    d.arrow("initial-execution",[(980,481),(980,550)])
    d.arrow("update-candidates",[(1360,481),(1360,550)])
    d.arrow("candidates-execution",[(1200,593),(1140,593)])
    d.arrow("joint-to-fixed",[(1360,237),(1360,310),(50,310),(50,443),(60,443)])
    d.arrow("joint-to-adjustable",[(810,310),(810,443),(820,443)])
    d.text("branch-repeat","按发布时刻重复更新；已执行时段始终冻结",820,645,720,30,14.25,"gray")
    d.group("settlement",40,750,1520,205,"共同结算与核验：交付后按附件四真实电价逐笔结算，库存跨日连续")
    for name,text,x in [
        ("actual-settlement","按交付时段真实电价结算\n固定计划核算普通费与紧急费\n可调整计划另逐笔计入调整净费",60),
        ("independent-audit","独立核验账单与物理约束\n能量守恒、库存边界与跨日连续\n逐笔计划重放及结果表回读",440),
        ("comparisons","对照与敏感性分析\n八种更新组合、价格点预测对照\n指定日期与离散精度检查",820),
        ("deliverables","固定计划与可调整计划结果表\n逐时轨迹、逐笔账本\n费用分解与方案对照",1200)]:
        d.node(name,text,x,825,"tall",fill="output" if x==1200 else None)
    for x in (380,760,1140):
        d.arrow("settlement-link",[(x,868),(x+60,868)])
    d.arrow("fixed-to-settlement",[(600,637),(600,715),(220,715),(220,825)])
    d.arrow("adjustable-to-settlement",[(980,637),(980,730),(220,730)])
    d.text("footer","未来真实电价不进入计划或候选评价；当前真实价在执行时揭示；两条策略分别维护自身库存",40,997,1520,28,15,"gray")
    d.save("问题四/第四问_技术路线图_参考风格.pptx")


if __name__ == "__main__":
    clean_existing()
    draw_q3()
    draw_q4()
