"""Maintain the appendix, source snapshot, package and project navigation."""
from pathlib import Path
from datetime import datetime
import ast,copy,hashlib,io,json,os,shutil,sys,tokenize,zipfile

ROOT=Path(__file__).resolve().parents[2]
PACKAGE=ROOT/'论文写作/附录与支撑材料_20260913'
SUPPORT=PACKAGE/'支撑材料'
WORK=ROOT/'tmp/materials_sync_20260913'
REFILL=ROOT/'tmp/result_template_refill_20260913'
NAMES=['result1.xlsx','result2.xlsx','result3.xlsx','result4-2.xlsx','result4-3.xlsx']
STEM='论文附录'


def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def core_excerpt(path,names):
    source=path.read_text(encoding='utf-8-sig');lines=source.splitlines();parts=[]
    for node in ast.parse(source).body:
        if isinstance(node,ast.FunctionDef) and node.name in names:
            raw='\n'.join(lines[node.lineno-1:node.end_lineno])
            tokens=[t for t in tokenize.generate_tokens(io.StringIO(raw).readline) if t.type!=tokenize.COMMENT]
            raw=tokenize.untokenize(tokens);fn=ast.parse(raw).body[0]
            if isinstance(fn.body[0],ast.Expr) and isinstance(fn.body[0].value,ast.Constant) and isinstance(fn.body[0].value.value,str):
                rows=raw.splitlines();del rows[fn.body[0].lineno-1:fn.body[0].end_lineno];raw='\n'.join(rows)
            parts.append('\n'.join(x.rstrip() for x in raw.splitlines() if x.strip()))
    assert len(parts)==len(names),(path,names)
    return '\n\n'.join(parts)


def build():
    from docx import Document
    from docx.shared import Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    WORK.mkdir(parents=True,exist_ok=True)
    baseline=PACKAGE/(STEM+'.docx')
    if not baseline.exists():baseline=PACKAGE/'论文尾部_附录修订版.docx'
    doc=Document(baseline)
    code_sample=next(p for p in doc.paragraphs if p.text.startswith('def solve('))
    code_properties=copy.deepcopy(code_sample._p.pPr)
    # Keep the supplied section geometry and document styles.
    for p in list(doc.paragraphs):p._element.getparent().remove(p._element)
    ai_text='本参赛队在竞赛过程中使用了 ChatGPT，主要用于代码检查与调试、基于 MATLAB 的科研绘图与结果可视化，以及时间标签核对、跨日补算实现、结果模板回填和支撑材料整理；参赛队已人工讨论并修改建模方案，并核验代码。详见支撑文件。'
    items=[('h2','AI工具使用声明'),('p',ai_text),('h1','附录'),('h2','A 支撑文件列表'),
      ('p','1. AI使用说明.docx：AI工具的使用范围与方式。'),
      ('p','2. 源码/代码/：四问模型、实验与绘图程序，以及结果导出和跨日补算程序。主版本与运行顺序见使用说明.md。'),
      ('p','3. 结果/：result1.xlsx、result2.xlsx、result3.xlsx、result4-2.xlsx、result4-3.xlsx；问题一至问题四子目录保留原模型的汇总和明细。'),
      ('p','4. 结果/跨日补算/：跨日预测记录、模板回填核验记录和时间口径说明。'),
      ('p','5. 使用说明.md：目录说明、复现顺序及结果时间口径。'),
      ('break',''),('h2','B 核心源码'),
      ('p','以下列出主要计算函数。代码依赖、参数和调用入口见对应完整源码；原模型与新增跨日处理分开保存。')]
    specs=[('问题一','代码/q1_solution.py',['solve']),
      ('问题二','代码/实验/q2_fresh_20260912/fresh_engine.py',['plan_lp','action']),
      ('问题三','代码/实验/q3_revision_a_20260912/q3_engine.py',['bellman_batch']),
      ('问题四','代码/实验/q4_independent_codex_20260912/model.py',['bellman']),
      ('时间对齐与跨日延拓','代码/结果导出/boundary_rules.py',['solve_next_cycle','extend_pv_hours','match_template_day'])]
    excerpts=[]
    for i,(title,rel,names) in enumerate(specs):
        if i in [2,4]:items.append(('break',''))
        code=core_excerpt(ROOT/rel,names)
        items.extend([('h3',title),('p','源码位置：源码/'+rel),('code',code)])
        excerpts.append(dict(section=title,source=rel,functions=names,source_sha256=digest(ROOT/rel),code=code))
    b=json.loads((REFILL/'boundary_prediction.json').read_text(encoding='utf-8'))
    items.extend([('h2','C 时间口径与跨日计算'),
      ('p','输入时标按十分钟区间的结束时刻解释，例如0:10表示00:00—00:10。结果模板时间栏保持原样，首段00:10—00:20对应模型第二段，末段对应次日00:00—00:10；末段不能直接用当天第一段替代。'),
      ('p',f'第一问的新增边界假设是光伏预测曲线延续一个日周期，电价和负载按题设每日重复。调用原求解器重新计算下一周期，并校验其期初储电量与前一期末相接。次日首段购电量为{b["q1"]["first_slot_kwh"]:.4f} kWh，储电衔接值为{b["q1"]["day_initial_kwh"]:.4f} kWh；该值依赖日周期延拓假设。'),
      ('p',f'第二问与4-2沿用原预测和优化方法向后补算，2026年1月1日首段为{b["first_slot_kwh"]:.4f} kWh。第三问与4-3沿用2025年12月31日18点预报覆盖次日1—18点，对19—24点采用此前7天零点发布的同一时刻预报均值延拓，再执行原插值、历史残差校准及优化，首段为{b["q3"]["first_slot_kwh"]:.4f} kWh。延拓预报不是附件中实际发布的新预报，也不是实际观测。'),
      ('p','购电表合计覆盖模板展示的00:10—次日00:10，充放电和紧急购电仍按各自模板指定的自然日统计；原论文全年费用对应自然日范围，二者不能直接混用。每行末段接入次日零点产生的计划，不将其表述为前一日零点已知的决策。4-2和4-3年末跨日费用采用预测电价。'),
      ('p','完整跨日计算见源码/代码/结果导出/boundary_prediction.py；模板数据准备、回填和验证分别见prepare_templates.py、refill_templates.mjs与verify_templates.py。')])
    markdown=[]
    for kind,text in items:
        if kind=='break':doc.add_page_break();continue
        if kind.startswith('h'):
            level=int(kind[1]);p=doc.add_heading(text,level)
            for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)
            markdown.extend(['#'*level+' '+text,''])
        elif kind=='code':
            for line in text.splitlines():
                p=doc.add_paragraph()
                if code_properties is not None:p._p.insert(0,copy.deepcopy(code_properties))
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=Pt(10)
                run=p.add_run(line);run.font.name='Consolas';run.font.size=Pt(8)
                fonts=run._element.get_or_add_rPr().rFonts
                fonts.set(qn('w:eastAsia'),'宋体')
            markdown.extend(['```python',text,'```',''])
        else:doc.add_paragraph(text);markdown.extend([text,''])
    doc.save(WORK/(STEM+'.docx'))
    (WORK/(STEM+'.md')).write_text('\n'.join(markdown),encoding='utf-8')
    ai=Document(SUPPORT/'AI使用说明.docx')
    ai.paragraphs[2].text='本参赛队在竞赛过程中使用 ChatGPT（OpenAI）辅助开展代码检查与调试、基于 MATLAB 的科研绘图与结果可视化，以及时间标签核对、跨日补算实现、结果模板回填和支撑材料整理。'
    ai.paragraphs[6].insert_paragraph_before('在结果整理环节，借助 ChatGPT 核对输入时标与结果模板的时段对应关系，实现并检查经参赛队选择的跨日预测或延拓假设，回填五份结果工作簿，并同步程序、附录和支撑材料。新增假设及数值来源在附录和跨日补算记录中单独说明。')
    ai.save(WORK/'AI使用说明.docx')
    check=Document(WORK/(STEM+'.docx'))
    actual='\n'.join(p.text for p in check.paragraphs)
    for e in excerpts:assert e['code'] in actual,e['source']
    assert len(check.tables)==0
    (WORK/'document_content_checks.json').write_text(json.dumps(dict(status='content_verified_render_pending',excerpts=excerpts,paragraphs=len(check.paragraphs),tables=0,baseline=str(baseline)),ensure_ascii=False,indent=2),encoding='utf-8')
    # This PDF is an independent layout of the same items, not a Word export.
    render_pdf(items,WORK/(STEM+'.pdf'))
    ai_items=[('h1' if p.style.name=='Heading 1' else 'h2' if p.style.name=='Heading 2' else 'p',p.text) for p in ai.paragraphs if p.text]
    render_pdf(ai_items,WORK/'AI使用说明_核对.pdf')
    print(json.dumps({'docx_built':2,'appendix_paragraphs':len(check.paragraphs),'core_sources':[e['source'] for e in excerpts],'pdf_renderer':'independent ReportLab layout; not a DOCX render'},ensure_ascii=False))


def render_pdf(items,path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Preformatted,KeepTogether
    from xml.sax.saxutils import escape
    pdfmetrics.registerFont(TTFont('Song','C:/Windows/Fonts/simsun.ttc',subfontIndex=0))
    pdfmetrics.registerFont(TTFont('Hei','C:/Windows/Fonts/simhei.ttf'))
    pdfmetrics.registerFont(TTFont('Code','C:/Windows/Fonts/consola.ttf'))
    body=ParagraphStyle('body',fontName='Song',fontSize=10.5,leading=16,spaceAfter=6)
    headings={k:ParagraphStyle(k,fontName='Hei',fontSize=size,leading=size+5,spaceBefore=8,spaceAfter=6,keepWithNext=True) for k,size in [('h1',15),('h2',12),('h3',11)]}
    code=ParagraphStyle('code',fontName='Code',fontSize=7.6,leading=10,spaceAfter=0)
    story=[]
    for kind,text in items:
        if kind=='break':story.append(PageBreak())
        elif kind=='code':
            for line in text.splitlines():
                while pdfmetrics.stringWidth(line,'Code',7.6)>430:
                    n=len(line)
                    while pdfmetrics.stringWidth(line[:n],'Code',7.6)>430:n-=1
                    split=line.rfind(' ',0,n)
                    if split<max(10,n//2):split=n
                    story.append(Preformatted(line[:split],code));line='    '+line[split:].lstrip()
                story.append(Preformatted(line or ' ',code))
        else:story.append(Paragraph(escape(text),headings.get(kind,body)))
    SimpleDocTemplate(str(path),pagesize=A4,leftMargin=70,rightMargin=70,topMargin=50,bottomMargin=50).build(story)


def sync_files():
    WORK.mkdir(parents=True,exist_ok=True)
    backup=ROOT/'论文写作/历史版本'/('整理前备份_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    backup.mkdir(parents=True,exist_ok=False)
    before={}
    def write_target(src,target):
        assert target.resolve().is_relative_to(ROOT.resolve())
        if target.exists() and digest(src)!=digest(target):
            dest=backup/target.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(target,dest)
        target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,target)
        assert digest(src)==digest(target)
    archive=PACKAGE/'支撑材料.zip'
    old_archive=backup/'支撑材料_更新前.zip';shutil.copy2(archive,old_archive)
    records=SUPPORT/'结果/跨日补算';records.mkdir(parents=True,exist_ok=True)
    for source,name in [(REFILL/'boundary_prediction.json','年末预测记录.json'),(REFILL/'verification.json','核验记录.json')]:write_target(source,records/name)
    old_note=ROOT/'论文写作/原模板重填_20260913/请先阅读_重填状态.txt'
    if old_note.exists():write_target(old_note,records/'时间口径说明.txt')
    for name in NAMES:
        source=REFILL/'filled'/name
        assert digest(source)==next(x['sha256'] for x in json.loads((REFILL/'verification.json').read_text(encoding='utf-8'))['verified'] if x['file']==name)
        write_target(source,SUPPORT/'结果'/name)
    copied={}
    for folder,dirs,files in os.walk(ROOT/'代码',followlinks=False):
        dirs[:]=[d for d in dirs if d not in ['node_modules','.deps','__pycache__','.git'] and not (getattr((Path(folder)/d).stat(follow_symlinks=False),'st_file_attributes',0)&1024)]
        for name in files:
            source=Path(folder)/name
            if source.suffix.lower() not in ['.py','.m','.ps1','.mjs','.js']:continue
            rel=source.relative_to(ROOT)
            target=SUPPORT/'源码'/rel
            write_target(source,target);copied[str(rel)]=digest(source)
    if '--include-documents' in sys.argv:
        for name in [STEM+'.docx',STEM+'.pdf',STEM+'.md']:
            write_target(WORK/name,PACKAGE/name)
        write_target(WORK/'AI使用说明.docx',SUPPORT/'AI使用说明.docx')
    (WORK/'sync_state.json').write_text(json.dumps(dict(backup=str(backup),source_hashes=copied),ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(synced_source_files=len(copied),backup=str(backup)),ensure_ascii=False))


def pack():
    state=json.loads((WORK/'sync_state.json').read_text(encoding='utf-8'))
    sources=state['source_hashes']
    for rel,h in sources.items():assert digest(ROOT/rel)==digest(SUPPORT/'源码'/rel)==h
    entries={p.relative_to(SUPPORT).as_posix():p for p in SUPPORT.rglob('*') if p.is_file() and not p.name.startswith('~$')}
    target=WORK/'支撑材料_同步.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for name,p in sorted(entries.items()):z.write(p,name)
    with zipfile.ZipFile(target) as z:
        assert set(z.namelist())==set(entries) and z.testzip() is None
        for name,p in entries.items():assert hashlib.sha256(z.read(name)).hexdigest()==digest(p)
    shutil.copy2(target,PACKAGE/'支撑材料.zip')
    from pypdf import PdfReader
    appendix=PACKAGE/(STEM+'.pdf')
    audit=PACKAGE/'核验';audit.mkdir(exist_ok=True)
    size=dict(limit_bytes=20000000,support_zip_bytes=target.stat().st_size,appendix_pdf_bytes=appendix.stat().st_size,
      appendix_pdf_pages=len(PdfReader(appendix).pages),remaining_full_paper_pdf_bytes=20000000-target.stat().st_size,
      remaining_body_without_appendix_bytes=20000000-target.stat().st_size-appendix.stat().st_size)
    (audit/'体积核验.json').write_text(json.dumps(size,ensure_ascii=False,indent=2),encoding='utf-8')
    doc_checks=json.loads((WORK/'document_content_checks.json').read_text(encoding='utf-8'))
    report=dict(status='synchronized',archive_files=len(entries),source_files=len(sources),
        archive_sha256=digest(PACKAGE/'支撑材料.zip'),source_hashes=sources,
        code_excerpts_verified=[{k:v for k,v in e.items() if k!='code'} for e in doc_checks['excerpts']],
        word_render_status='not_verified',pdf_render_status='independent layout verified separately',
        result_hashes={n:digest(SUPPORT/'结果'/n) for n in NAMES},backup=state['backup'])
    (audit/'交付核验记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(archive_files=len(entries),source_files=len(sources),size=size),ensure_ascii=False))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    {'build':build,'sync':sync_files,'pack':pack}[sys.argv[1]]()
