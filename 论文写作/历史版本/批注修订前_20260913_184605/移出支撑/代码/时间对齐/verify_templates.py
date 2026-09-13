"""Independent read-back of every filled cell and untouched template header."""
from pathlib import Path
from datetime import datetime,time
import hashlib,json,sys
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'tmp/result_template_refill_20260913'


def normalized(value):
    if isinstance(value,datetime):return to_excel(value)
    if isinstance(value,time):return (value.hour*3600+value.minute*60+value.second)/86400
    return value


def compare(ws,row0,col0,expected):
    errors=[]
    for i,row in enumerate(expected,row0):
        for j,value in enumerate(row,col0):
            actual=normalized(ws.cell(i,j).value)
            if isinstance(value,(int,float)):
                assert isinstance(actual,(int,float)),(ws.title,i,j,actual,value)
                errors.append(abs(actual-value))
            else:assert actual==value,(ws.title,i,j,actual,value)
    error=max(errors,default=0.)
    assert error<1e-8,(ws.title,error)
    return sum(map(len,expected)),error


def main():
    payload=json.loads((WORK/'payload.json').read_text(encoding='utf-8'))
    checks=[]
    for spec in payload['workbooks']:
        if not spec['ready']:continue
        name=spec['file'];file=WORK/'filled'/name
        original=load_workbook(ROOT/'基础数据'/name,data_only=False)
        saved=load_workbook(file,data_only=False)
        assert saved.sheetnames==original.sheetnames
        cells=0;error=0.
        for sheet,values in spec['plans'].items():
            ws=saved[sheet];template=original[sheet]
            if spec['kind']=='single':
                assert [ws.cell(i,1).value for i in range(1,146)]==[template.cell(i,1).value for i in range(1,146)]
            else:
                assert [ws.cell(1,j).value for j in range(1,148)]==[template.cell(1,j).value for j in range(1,148)]
                assert [normalized(ws.cell(i,1).value) for i in range(2,336)]==[normalized(template.cell(i,1).value) for i in range(2,336)]
                for row in values:assert abs(sum(row[:144])-row[144])<1e-6
            count,err=compare(ws,2,2,values);cells+=count;error=max(error,err)
        store=saved['充放电量'];src=original['充放电量']
        assert [store.cell(1,j).value for j in range(1,src.max_column+1)]==[src.cell(1,j).value for j in range(1,src.max_column+1)]
        if spec['kind']=='single':
            count,err=compare(store,2,2,spec['storage']);cells+=count;error=max(error,err)
            count,err=compare(store,2,5,[[spec['start']],[spec['end']]]);cells+=count;error=max(error,err)
        else:
            count,err=compare(store,2,1,spec['storage']);cells+=count;error=max(error,err)
            em=saved['紧急购电量'];template=original['紧急购电量']
            assert [em.cell(1,j).value for j in range(1,4)]==[template.cell(1,j).value for j in range(1,4)]
            count,err=compare(em,2,1,spec['emergency']);cells+=count;error=max(error,err)
        for ws in saved:
            assert not any(c.data_type in ('e','f') for row in ws for c in row)
        checks.append(dict(file=name,cells_compared=cells,max_absolute_error=error,
                           template_sheet_names_and_headers_preserved=True,
                           full_date_sequence_preserved=True,formula_errors=0,
                           sha256=hashlib.sha256(file.read_bytes()).hexdigest()))
        original.close();saved.close()
    for name,expected in payload['source_hashes'].items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==expected,('Source changed',name)
    report=dict(status='complete_with_authorized_boundary_assumptions' if all(s['ready'] for s in payload['workbooks']) else 'partial_pending_boundary_resolution',
                verified=checks,original_sources_unchanged=True,
                waiting=[x['file'] for x in payload['workbooks'] if not x['ready']],
                missing_inputs=payload['missing_inputs'],
                q1_status='Fresh next-day solve with explicitly authorized repeating PV forecast.' if payload['workbooks'][0]['ready'] else 'Next-day calculation not ready.',
                time_mapping='Input 00:10 ends 00:00-00:10. Original template preserved; values matched by actual delivery date/time.',
                annual_table_totals='Sum of the displayed 00:10 through next-day 00:10 window; storage and emergency tables retain their stated natural-day intervals.',
                boundary_prediction='Q1 repeats the supplied PV forecast and solves the next daily cycle. Q3 and Q4-3 use the authorized extension of the last available PV forecast. Q2 and Q4-2 use the original prediction methods. Original optimization sources are unchanged. Year-end Q4 costs use forecast prices.',
                issue_time_caveat='Next-day first slot is from the next-day model plan; it is not represented as a plan known at the previous midnight.',
                original_source_hashes=payload['source_hashes'])
    (WORK/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['status','verified','original_sources_unchanged','waiting']},ensure_ascii=False))


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8');main()
