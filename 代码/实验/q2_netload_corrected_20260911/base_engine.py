"""Independent Q2: quantile net-load forecast and causal Markov storage control.
All mutations confined to this script's directory. Original code is never imported.
"""
from pathlib import Path
import sys, os, json, csv, hashlib, time, platform, argparse
from datetime import datetime, timedelta
import numpy as np
import scipy
from scipy.optimize import linprog
from scipy.sparse import lil_matrix
from numpy.lib.stride_tricks import sliding_window_view
import lightgbm as lgb

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
sys.path.append(str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/Lib/site-packages'))
import openpyxl
OUT = BASE/'results'
OUT.mkdir(exist_ok=True)
N=144; DT=1/6; ETA=.9; CAP=5000/6; EMIN=1200.; EMAX=10800.
QUANTILES=np.array([.2,.5,.65,.8,.9]); SEED=20260911
GRID=np.arange(EMIN, EMAX+1,60.)
TOL=2e-5

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,x): Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def date(d): return (datetime(2025,1,1)+timedelta(days=int(d))).strftime('%Y-%m-%d')
def csvout(p,heads,rows):
    with Path(p).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(heads);w.writerows(rows)
def log(*args): print(datetime.now().isoformat(timespec='seconds'),*args,flush=True)

def load():
    wb=openpyxl.load_workbook(ROOT/'基础数据/附件2.xlsx',read_only=True,data_only=True)
    ys=[]; heads=[]
    for ws in wb:
        rows=list(ws.values)
        assert len(rows)==366 and len(rows[0])==145
        heads.append(rows[0][1:])
        assert all(r[0]==datetime(2025,1,1)+timedelta(days=d) for d,r in enumerate(rows[1:]))
        ys.append(np.array([r[1:] for r in rows[1:]],float))
    wb.close()
    assert heads[0]==heads[1] and len(ys)==2
    y=np.asarray(ys);assert y.shape==(2,365,N) and np.isfinite(y).all() and y.min()>=0
    wb=openpyxl.load_workbook(ROOT/'基础数据/附件1.xlsx',read_only=True,data_only=True)
    rows=list(wb.active.values);wb.close()
    def minute(x):
        if hasattr(x,'hour'):return x.hour*60+x.minute
        if x=='0:00+1':return 1440
        h,m=map(int,str(x).split(':'));return h*60+m
    assert [minute(r[0]) for r in rows[1:]]==[minute(x) for x in heads[0]]==list(range(10,1441,10))
    p=np.array([r[1] for r in rows[1:]],float);assert p.shape==(144,) and p.min()>0
    dump(OUT/'data_audit.json',{'shape':list(y.shape),'min':y.min(axis=(1,2)).tolist(),'max':y.max(axis=(1,2)).tolist(),
        'dates':[date(0),date(364)],'interval_convention':'end timestamps, first 00:00-00:10, last 23:50-24:00'})
    return y,p

def features(y,d):
    """All rows of day d use exclusively days < d, including training rows."""
    assert d>=1
    t=np.arange(N)/N;wd=(datetime(2025,1,1)+timedelta(days=d)).weekday()
    cols=[t,np.sin(2*np.pi*t),np.cos(2*np.pi*t),np.sin(4*np.pi*t),np.cos(4*np.pi*t)]
    cols += [np.full(N,x) for x in [wd,wd>=5,np.sin(2*np.pi*wd/7),np.cos(2*np.pi*wd/7)]]
    ids=[j for j in range(max(0,d-35),d) if (d-j)%7==0]
    if not ids: ids=list(range(max(0,d-7),d))
    for a in y:
        hist=a[max(0,d-7):d]
        cols.extend([a[d-1],a[max(0,d-7)],a[max(0,d-3):d].mean(axis=0),hist.mean(axis=0),hist.std(axis=0),a[ids].mean(axis=0)])
        cols.extend([np.full(N,a[d-1].mean()),np.full(N,hist.mean())])
    return np.array(cols,dtype=np.float64).T

def one_fit(X,Z,x,days,alpha):
    params=dict(objective='quantile',alpha=float(alpha),verbosity=-1,num_threads=2,seed=SEED,
        deterministic=True,force_col_wise=True,num_leaves=15,max_depth=4,learning_rate=.05,
        min_data_in_leaf=48,lambda_l2=5.,feature_fraction=1.,bagging_fraction=1.)
    w=np.repeat(2.**(-(days[-1]-days)/28),N)
    model=lgb.train(params,lgb.Dataset(X,label=Z/1000,weight=w),num_boost_round=120)
    return model.predict(x,num_threads=2)*1000

def forecast_day(y,d,Xall):
    ids=np.arange(max(7,d-56),d)
    X=Xall[ids].reshape(-1,Xall.shape[-1]);x=Xall[d]
    net=y[0]-y[1]
    direct=np.array([one_fit(X,net[ids].ravel(),x,ids,a) for a in QUANTILES])
    separate=np.array([np.maximum(one_fit(X,y[v,ids].ravel(),x,ids,.5),0) for v in range(2)])
    return np.sort(direct,axis=0),separate

def forecasts(y):
    folder=OUT/'forecast_days';folder.mkdir(exist_ok=True)
    X=np.zeros((365,N,25))
    for d in range(1,365): X[d]=features(y,d)
    net=y[0]-y[1]
    base=np.zeros((365,5,N));pred=np.zeros_like(base);separate=np.zeros((365,2,N))
    direct=np.zeros((365,5,N));family='direct';selection={}
    for d in range(1,365):
        if d<14:
            direct[d]=net[d-1][None,:];separate[d]=y[:,d-1]
        else:
            f=folder/f'{date(d)}.npz'
            if f.exists():
                z=np.load(f);direct[d]=z['direct'];separate[d]=z['separate']
            else:
                direct[d],separate[d]=forecast_day(y,d,X)
                np.savez_compressed(f,direct=direct[d],separate=separate[d])
        if d==21:
            a=float(abs(direct[14:21,1]-net[14:21]).mean())
            b=float(abs(separate[14:21,0]-separate[14:21,1]-net[14:21]).mean())
            family='direct' if a<=b else 'separate_center'
            selection={'validation':'Jan 15-21','frozen_at':'Jan 22 00:00','direct_MAE':a,'separate_MAE':b,'family':family}
            log('forecast family',selection)
        base[d]=direct[d]
        if family=='separate_center':base[d]+=separate[d,0]-separate[d,1]-direct[d,1]
        pred[d]=base[d]
        ids=np.arange(max(14,d-30),d)
        if len(ids)>=4:
            for block in range(4):
                sl=slice(block*36,(block+1)*36)
                for k,a in enumerate(QUANTILES):
                    errors=net[ids,sl]-base[ids,k,sl]
                    pred[d,k,sl]+=np.quantile(errors,float(a))
            pred[d]=np.sort(pred[d],axis=0)
        if d%15==0:log('forecast',date(d))
    np.savez_compressed(OUT/'forecasts.npz',base=base,pred=pred,direct=direct,separate=separate)
    dump(OUT/'forecast_selection.json',selection)
    return pred

def markov(y,pred,d):
    """Only preceding dates enter this probability model."""
    ids=np.arange(max(14,d-30),d);assert len(ids)>0
    real=(y[0]-y[1])[ids]
    scale_hist=np.maximum((pred[ids,3]-pred[ids,0])/1.683,100.)
    z=(real-pred[ids,1])/scale_hist
    states=np.digitize(z,[-.84,.84])
    scale=np.maximum((pred[d,3]-pred[d,0])/1.683,100.)
    means=np.empty((N,3));P=np.empty((N-1,3,3))
    count=np.ones((3,3))
    np.add.at(count,(states[:,:-1].ravel(),states[:,1:].ravel()),1)
    global_P=count/count.sum(axis=1,keepdims=True)
    fallback=np.array([-1.4,0.,1.4])
    for t in range(N):
        a=max(0,t-3);b=min(N,t+4);zs=z[:,a:b];ss=states[:,a:b]
        for s in range(3):
            vals=zs[ss==s]
            means[t,s]=(vals.sum()+2*fallback[s])/(len(vals)+2)
        if t<N-1:
            a=max(0,t-3);b=min(N-1,t+4)
            cnt=3*global_P.copy()
            np.add.at(cnt,(states[:,a:b].ravel(),states[:,a+1:b+1].ravel()),1)
            P[t]=cnt/cnt.sum(axis=1,keepdims=True)
    initial=np.bincount(states[:,0],minlength=3)+np.array([.6,1.8,.6])
    initial=initial/initial.sum()
    net=(pred[d,1,:,None]+scale[:,None]*means)*DT
    assert np.isfinite(net).all() and np.allclose(P.sum(axis=2),1)
    return dict(net=net,P=P,initial=initial,scale=scale,center=pred[d,1])

def plan_lp(p,net,E0,kappa):
    """Deterministic forecast LP only proposes g; recourse is never executed."""
    H=len(p);dim=5*H+1;g=0;c=H;q=2*H;w=3*H;e=4*H
    obj=np.zeros(dim);obj[:H]=p;obj[-1]=-kappa
    bounds=[(0,None)]*dim
    A=lil_matrix((2*H,dim));rhs=np.r_[net,np.zeros(H)]
    for t in range(H):
        A[t,[g+t,c+t,q+t,w+t]]=[1,-1,1,-1]
        A[H+t,[e+t+1,e+t,c+t,q+t]]=[1,-1,-ETA,1/ETA]
        bounds[c+t]=bounds[q+t]=(0,CAP)
    for t in range(H+1):bounds[e+t]=(EMIN,EMAX)
    bounds[e]=(E0,E0)
    res=linprog(obj,A_eq=A.tocsr(),b_eq=rhs,bounds=bounds,method='highs')
    assert res.success,res.message
    assert abs(A@res.x-rhs).max()<TOL
    return np.maximum(res.x[:H],0)

def grid_value(v,r,p,grid):
    """Exact minimization of the linearly interpolated continuation on a grid.
    Actual actions respect direction; surplus charging is weakly dominant.
    """
    if r<=0:
        end=np.minimum(grid+ETA*min(-r,CAP),EMAX)
        return np.interp(end,grid,v)
    distance=min(r/ETA,CAP/ETA)
    lo=np.maximum(EMIN,grid-distance)
    width=min(len(grid),int(np.floor(distance/(grid[1]-grid[0])+1e-10))+1)
    width=min(width,len(grid))
    x=v+5*p*ETA*grid
    # Window consists of current grid point and preceding width-1 points.
    pad=np.pad(x,(width-1,0),constant_values=np.inf)
    knot_min=sliding_window_view(pad,width).min(axis=1)
    boundary=np.interp(lo,grid,v)+5*p*ETA*lo
    return 5*p*(r-ETA*grid)+np.minimum(knot_min,boundary)

def bellman(p,g,m,kappa,grid=GRID):
    H=len(p);V=np.empty((H+1,3,len(grid)));V[H]=-kappa*grid
    for t in range(H-1,-1,-1):
        for s in range(3):
            future=V[t+1,s] if t==H-1 else m['P'][t,s]@V[t+1]
            V[t,s]=grid_value(future,m['net'][t,s]-g[t],p[t],grid)
    assert np.isfinite(V).all()
    # More free initial stored energy cannot increase minimum cost.
    assert np.diff(V,axis=2).max()<TOL
    return V

def choose_plan(p,pred_day,m,E0,kappa,median_only=False,grid=GRID):
    marginal=m['initial'].copy();mean=[]
    for t in range(N):
        mean.append(marginal@m['net'][t])
        if t<N-1:marginal=marginal@m['P'][t]
    nets=[pred_day[k]*DT for k in [1,2,3,4]]+[np.asarray(mean)]
    labels=['q50','q65','q80','q90','markov_mean']
    if median_only:nets=nets[:1];labels=labels[:1]
    candidates=[];best=None
    for label,net in zip(labels,nets):
        g=plan_lp(p,net,E0,kappa)
        V=bellman(p,g,m,kappa,grid)
        score=float(p@g+sum(m['initial'][s]*np.interp(E0,grid,V[0,s]) for s in range(3)))
        candidates.append({'candidate':label,'expected_planning_loss':score,'normal_cost':float(p@g)})
        if best is None or score<best[0]-1e-7:best=(score,label,g,V)
    return best[2],best[3],best[1],candidates

def action(E,r,p,future,grid=GRID):
    if r<=0:
        c=max(0,min(-r,CAP,(EMAX-E)/ETA));return c,0.,0.,max(0,-r-c),E+ETA*c
    lo=max(EMIN,E-min(r,CAP)/ETA)
    knots=grid[(grid>=lo-1e-9)&(grid<=E+1e-9)]
    knots=np.unique(np.r_[lo,E,knots]);knots=np.clip(knots,lo,E)
    costs=5*p*np.maximum(r-ETA*(E-knots),0)+np.interp(knots,grid,future)
    nxt=float(knots[np.argmin(costs)])
    q=max(0,ETA*(E-nxt));e=max(0,r-q)
    return 0.,q,e,0.,nxt

def execute(actual,p,g,E0,m,V,grid=GRID,greedy=False):
    rows=[];E=float(E0)
    for t in range(N):
        net=float(actual[t])*DT;r=net-g[t]
        if greedy:
            if r<=0:
                c=max(0,min(-r,CAP,(EMAX-E)/ETA));q=e=0.;w=max(0,-r-c)
            else:
                q=max(0,min(r,CAP,(E-EMIN)*ETA));c=w=0.;e=max(0,r-q)
            nxt=E+ETA*c-q/ETA
        else:
            s=int(np.digitize((actual[t]-m['center'][t])/m['scale'][t],[-.84,.84]))
            future=V[t+1,s] if t==N-1 else m['P'][t,s]@V[t+1]
            c,q,e,w,nxt=action(E,r,p[t],future,grid)
        rows.append([g[t],c,q,e,w,E,nxt]);E=nxt
    arr=np.asarray(rows);validate(arr,actual,p)
    return arr

def validate(arr,actual,p):
    a=np.asarray(arr);g,c,q,e,w,before,after=a.T
    errs={'balance':float(abs(g+e+q-c-w-actual*DT).max()),
        'state':float(abs(after-before-ETA*c+q/ETA).max()),
        'continuity':float(abs(before[1:]-after[:-1]).max()) if len(a)>1 else 0.,
        'simultaneous':float(np.minimum(c,q).max()),
        'soc':float(max(0,EMIN-before.min(),EMIN-after.min(),before.max()-EMAX,after.max()-EMAX)),
        'power':float(max(0,c.max()-CAP,q.max()-CAP)),
        'nonnegative':float(max(0,-a[:,:5].min())),
        'emergency_charging':float(np.minimum(c,e).max())}
    assert np.isfinite(a).all() and max(errs.values())<TOL,errs
    return errs

def warm(y):
    net=y[0]-y[1];E=6000.;days=[]
    for d in range(21):
        g=np.zeros(N) if d==0 else np.maximum(net[d-1]*DT,0)
        a=execute(net[d],np.ones(N),g,E,None,None,greedy=True)
        days.append(a);E=a[-1,-1]
    return np.asarray(days)

def cost(a,p): return float(np.sum(a[...,0]*p)+5*np.sum(a[...,3]*p))

def backtest(y,p,pred):
    warmup=warm(y);np.save(OUT/'january_warmup.npy',warmup)
    kappas=[float(.9*p.mean()),float(p.mean()/ETA)]
    val=[]
    for k in kappas:
        E=warmup[-1,-1,-1];days=[]
        for d in range(21,31):
            m=markov(y,pred,d);g,V,label,cands=choose_plan(p,pred[d],m,E,k)
            a=execute((y[0]-y[1])[d],p,g,E,m,V);E=a[-1,-1];days.append(a)
        days=np.asarray(days);np.save(OUT/f'validation_{k:.9f}.npy',days)
        val.append({'kappa':k,'cost':cost(days,p),'emergency_kWh':float(days[:,:,3].sum()),'E_feb1':float(E)})
        log('January validation',val[-1])
    selected=min(val,key=lambda x:x['cost']);k=selected['kappa'];E0=selected['E_feb1']
    dump(OUT/'controller_selection.json',{'candidates':val,'selected':selected,'period':'Jan22-31','frozen_at':'Feb1 00:00'})
    state={s:E0 for s in ['main','median_only','greedy_same_plan']}
    runs={s:[] for s in state}
    folder=OUT/'days';folder.mkdir(exist_ok=True)
    for d in range(31,365):
        path=folder/f'{date(d)}.npz';meta_path=folder/f'{date(d)}.json'
        if path.exists() and meta_path.exists():
            z=np.load(path)
            for s in state:
                a=z[s];assert abs(a[0,5]-state[s])<TOL
                validate(a,(y[0]-y[1])[d],p);runs[s].append(a);state[s]=float(a[-1,-1])
        else:
            m=markov(y,pred,d)
            g,V,label,cands=choose_plan(p,pred[d],m,state['main'],k)
            a=execute((y[0]-y[1])[d],p,g,state['main'],m,V)
            gm,Vm,_,_=choose_plan(p,pred[d],m,state['median_only'],k,median_only=True)
            am=execute((y[0]-y[1])[d],p,gm,state['median_only'],m,Vm)
            ag=execute((y[0]-y[1])[d],p,g,state['greedy_same_plan'],m,V,greedy=True)
            daily={'main':a,'median_only':am,'greedy_same_plan':ag}
            np.savez_compressed(path,**daily)
            dump(meta_path,{'date':date(d),'chosen_candidate':label,'candidates':cands})
            for s,a in daily.items():runs[s].append(a);state[s]=float(a[-1,-1])
        if (d-31)%10==0 or d==364:log('backtest',date(d),{s:round(cost(np.array(runs[s]),p),2) for s in state})
    for s in runs:runs[s]=np.asarray(runs[s]);np.save(OUT/f'{s}.npy',runs[s])
    return runs,k

def metrics(a,p):
    daily=np.sum(a[:,:,0]*p+5*a[:,:,3]*p,axis=1)
    return {'normal_cost':float(np.sum(a[:,:,0]*p)),'emergency_cost':float(5*np.sum(a[:,:,3]*p)),
        'total_cost':float(daily.sum()),'planned_kWh':float(a[:,:,0].sum()),'emergency_kWh':float(a[:,:,3].sum()),
        'spill_kWh':float(a[:,:,4].sum()),'charge_kWh':float(a[:,:,1].sum()),'discharge_kWh':float(a[:,:,2].sum()),
        'emergency_days':int(np.sum(np.any(a[:,:,3]>TOL,axis=1))),
        'emergency_intervals':int(np.sum(a[:,:,3]>TOL)),'E_start':float(a[0,0,5]),'E_end':float(a[-1,-1,-1]),
        'max_daily_cost':float(daily.max()),'top_5pct_daily_mean':float(np.sort(daily)[-int(np.ceil(.05*len(daily))):].mean())}

def verify_algorithms():
    rng=np.random.default_rng(SEED)
    tests=0;err=0.
    for r in [-1800.,-60.,0.,30.,450.,1600.]:
        v=-.55*GRID+180*np.exp(-(GRID-EMIN)/1500)
        vals=grid_value(v,r,.8,GRID)
        for j in rng.choice(len(GRID),8,replace=False):
            c,q,e,w,en=action(GRID[j],r,.8,v)
            exact=4*e+np.interp(en,GRID,v)
            err=max(err,abs(exact-vals[j]));tests+=1
    assert err<1e-7,err
    # E[min cost after seeing a future branch] must not replace min E[cost].
    p=np.array([1.,1.]);g=np.array([0.,0.]);gg=np.array([1200.,1800.,2400.])
    m={'net':np.array([[600.,600.,600.],[0.,900.,1800.]]),
       'P':np.array([[[.2,.5,.3],[.2,.5,.3],[.2,.5,.3]]])}
    V=bellman(p,g,m,.5,gg)
    for s in range(3):
        future=m['P'][0,s]@V[1]
        for j,E in enumerate(gg):
            c,q,e,w,en=action(E,600.,1.,future,gg)
            assert abs(5*e+np.interp(en,gg,future)-V[0,s,j])<1e-7
    return {'grid_operator_cases':tests,'max_error':err,'expectation_before_action_check':True}

def validate_results(y,p,pred,runs,k):
    checks={'algorithm':verify_algorithms(),'physical':{}}
    for s,a in runs.items():
        err=validate(a.reshape(-1,7),(y[0]-y[1])[31:].ravel(),np.tile(p,334))
        checks['physical'][s]=err
    d=31;m=markov(y,pred,d);E=runs['main'][0,0,5]
    g,V,_,_=choose_plan(p,pred[d],m,E,k)
    actual=(y[0]-y[1])[d].copy();mutated=actual.copy();mutated[72:]+=50000
    a=execute(actual,p,g,E,m,V);b=execute(mutated,p,g,E,m,V)
    assert np.array_equal(a[:72],b[:72])
    altered=y.copy();altered[:,d:]*=7
    assert np.array_equal(features(y,d),features(altered,d))
    mm=markov(altered,pred,d)
    for key in m:assert np.array_equal(m[key],mm[key])
    gg,_,_,_=choose_plan(p,pred[d],mm,E,k);assert np.array_equal(g,gg)
    checks['causality']={'first_72_steps_unchanged_after_future_actual_mutation':True,
        'dayahead_features_probabilities_and_plan_unchanged_after_current_and_future_days_mutation':True,
        'scope':'regression tests plus functions use only past days/current observation; not a universal statistical guarantee'}
    refine=[]
    for ds in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
        d=(datetime.fromisoformat(ds)-datetime(2025,1,1)).days
        a=runs['main'][d-31];m=markov(y,pred,d);small=np.arange(EMIN,EMAX+1,30.)
        v=bellman(p,a[:,0],m,k,small)
        b=execute((y[0]-y[1])[d],p,a[:,0],a[0,5],m,v,small)
        refine.append({'date':ds,'same_g_and_initial_E':True,'cost_grid60':cost(a,p),'cost_grid30':cost(b,p),
            'E_end_grid60':float(a[-1,-1]),'E_end_grid30':float(b[-1,-1]),'max_step_E_difference':float(abs(a[:,6]-b[:,6]).max())})
    checks['grid_refinement']=refine
    dump(OUT/'validation.json',checks)
    return checks

def check_preservation():
    originals=json.loads((BASE/'original_files_sha256.json').read_text(encoding='utf-8-sig'))
    changed=[]
    for item in originals:
        # 实验目录及其归档文档由本实验维护；保护清单只核对实验之外的用户文件。
        rel=item['path'].replace('/','\\')
        if rel.startswith('代码\\实验\\') or rel.startswith('建模思路\\问题二\\') or rel.startswith('论文写作\\问题二\\'):
            continue
        p=ROOT/item['path']
        if not p.is_file() or sha(p).lower()!=item['sha256'].lower():changed.append(item['path'])
    res={'checked_files':len(originals),'changed_or_missing':changed,'excluded':'git, skill metadata, old .deps, pycache, managed experiment/document folders'}
    dump(OUT/'original_file_preservation.json',res)
    return res

def export(y,p,pred,runs,k,checks):
    net=y[0]-y[1];summary={s:metrics(a,p) for s,a in runs.items()}
    forecast_scores=[]
    for q,a in enumerate(QUANTILES):
        err=net[31:]-pred[31:,q]
        forecast_scores.append({'quantile':float(a),'observed_fraction_below':float((err<=0).mean()),
            'pinball_kW':float(np.maximum(a*err,(a-1)*err).mean())})
    mae=float(abs(net[31:]-pred[31:,1]).mean());rmse=float(np.sqrt(((net[31:]-pred[31:,1])**2).mean()))
    old_path=ROOT/'代码/results/q2/summary.json'
    old=json.loads(old_path.read_text(encoding='utf-8')) if old_path.exists() else {}
    data={'period':'2025-02-01..2025-12-31','days':334,'intervals':334*N,'kappa':k,
        'new':summary,'forecast':{'net_MAE_kW':mae,'net_RMSE_kW':rmse,'quantiles':forecast_scores},
        'old_summary_readonly':old,'limitations':['finite purchase candidate set','first-order 3-state error approximation',
        '60kWh value grid; continuous actual E','emergency purchase cannot charge storage in this policy',
        'information measured within current interval; interval averages not known ahead of interval',
        'same existing year used for method development, not a pristine blind holdout']}
    dump(OUT/'summary.json',data)
    header=['日期','时段开始','时段结束','电价_元每kWh','实际净负荷_kW','计划购电_kWh','充电_kWh','放电_kWh','紧急购电_kWh','富余电量_kWh','期初储电量_kWh','期末储电量_kWh','计划费用_元','紧急费用_元']
    def clock(t):return f'{t//6:02d}:{t%6*10:02d}'
    daily=[]
    for s,a in runs.items():
        def rows():
            for di,d in enumerate(range(31,365)):
                for t in range(N):yield [date(d),clock(t),clock(t+1),p[t],net[d,t],*a[di,t],p[t]*a[di,t,0],5*p[t]*a[di,t,3]]
        csvout(OUT/f'{s}_full_strategy.csv',header,rows())
        for di,d in enumerate(range(31,365)):
            daily.append([s,date(d),float(p@a[di,:,0]),float(5*p@a[di,:,3]),cost(a[di],p),float(a[di,:,3].sum()),float(a[di,0,5]),float(a[di,-1,-1])])
    csvout(OUT/'daily_metrics.csv',['策略','日期','计划费用_元','紧急费用_元','总费用_元','紧急电量_kWh','日初电量_kWh','日末电量_kWh'],daily)
    csvout(OUT/'forecast_quantiles.csv',['日期','时段结束','实际净负荷_kW']+[f'q{int(x*100)}_kW' for x in QUANTILES],
        ([date(d),clock(t+1),net[d,t],*pred[d,:,t]] for d in range(31,365) for t in range(N)))
    wb=openpyxl.Workbook();ws=wb.active;ws.title='费用对照'
    ws.append(['策略']+list(summary['main']))
    for s,a in summary.items():ws.append([s]+list(a.values()))
    ws=wb.create_sheet('每日费用');ws.append(['策略','日期','计划费用_元','紧急费用_元','总费用_元','紧急电量_kWh','日初电量_kWh','日末电量_kWh'])
    for row in daily:ws.append(row)
    for ds in ['2025-03-20','2025-06-21','2025-09-23','2025-12-21']:
        d=(datetime.fromisoformat(ds)-datetime(2025,1,1)).days;ws=wb.create_sheet(ds);ws.append(header)
        for t in range(N):
            a=runs['main'][d-31,t];ws.append([ds,clock(t),clock(t+1),p[t],net[d,t],*a,p[t]*a[0],5*p[t]*a[3]])
    ws=wb.create_sheet('全年主方案');ws.append(header)
    for d in range(31,365):
        for t in range(N):
            a=runs['main'][d-31,t];ws.append([date(d),clock(t),clock(t+1),p[t],net[d,t],*a,p[t]*a[0],5*p[t]*a[3]])
    ws=wb.create_sheet('口径说明')
    for row in [['项目','口径'],['时间','结束标签，00:10 表示 00:00—00:10；非官方开始标签模板'],['状态','独立实验，待审核；不是自动提交版'],['初末库存','各策略见费用对照，跨日连续'],['近似','三状态 Markov、60kWh价值网格、五个计划候选'],['富余电量','包含付费计划电量和光伏，不能全部称为弃光'],['真实费用','正常按计划量结算，紧急按当时电价五倍；不包含终端价值']]:ws.append(row)
    from openpyxl.styles import Font,PatternFill
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='264653')
        for col in ws.columns:
            first=col[0];ws.column_dimensions[first.column_letter].width=min(28,max(16,len(str(first.value))+3))
    xlsx=OUT/'result2_独立实验审核版.xlsx';wb.save(xlsx);wb.close()
    read=openpyxl.load_workbook(xlsx,read_only=True,data_only=True)
    ws=read['全年主方案'];assert ws.max_row==334*N+1
    normal=emergency=ekwh=0.
    for row in ws.iter_rows(min_row=2,values_only=True):normal+=row[-2];emergency+=row[-1];ekwh+=row[8]
    read.close()
    assert abs(normal-summary['main']['normal_cost'])<1e-5 and abs(emergency-summary['main']['emergency_cost'])<1e-5
    assert abs(ekwh-summary['main']['emergency_kWh'])<1e-5
    checks['xlsx_readback']={'rows':334*N,'normal_cost':normal,'emergency_cost':emergency,'emergency_kWh':ekwh}
    dump(OUT/'validation.json',checks)
    make_plots(runs,p,pred,net)
    return data

def make_plots(runs,p,pred,net):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'Microsoft YaHei','font.size':10,'axes.unicode_minus':False,'svg.fonttype':'none'})
    figdir=BASE/'figures';figdir.mkdir(exist_ok=True)
    fig,ax=plt.subplots(figsize=(7,4))
    labels={'main':'概率购电＋随机控制','median_only':'中位数购电＋随机控制','greedy_same_plan':'相同计划＋简单控制'}
    for (s,a),color,style in zip(runs.items(),['#0072B2','#D55E00','#009E73'],['-','--',':']):
        daily=np.sum(a[:,:,0]*p+5*a[:,:,3]*p,axis=1)
        ax.plot(np.arange(1,335),daily.cumsum()/10000,label=labels[s],color=color,linestyle=style)
    ax.set(xlabel='自 2 月 1 日起的运行天数',ylabel='累计实际购电费用（万元）');ax.legend();fig.tight_layout()
    for ext in ['png','svg']:fig.savefig(figdir/f'累计费用对照.{ext}',dpi=300)
    plt.close(fig)
    d=78;t=np.arange(N)/6
    fig,ax=plt.subplots(figsize=(7,4))
    ax.fill_between(t,pred[d,0],pred[d,3],alpha=.2,color='#0072B2',label='q20—q80 区间')
    ax.plot(t,pred[d,1],color='#0072B2',label='中位数预测')
    ax.plot(t,net[d],color='#D55E00',linestyle='--',label='实际净负荷')
    ax.set(xlabel='时段开始时刻（小时）',ylabel='净负荷（kW）',xlim=(0,24));ax.legend();fig.tight_layout()
    for ext in ['png','svg']:fig.savefig(figdir/f'3月20日概率预测.{ext}',dpi=300)
    plt.close(fig)

def signature(y,p):
    return {'code_sha256':sha(__file__),'inputs':{str(path.relative_to(ROOT)):sha(path) for path in [ROOT/'基础数据/附件1.xlsx',ROOT/'基础数据/附件2.xlsx']},
        'numpy':np.__version__,'scipy':scipy.__version__,'lightgbm':lgb.__version__}

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--stage',choices=['all','forecast','smoke','run','export'],default='all');args=parser.parse_args()
    started=time.time();y,p=load();sig=signature(y,p);sigpath=OUT/'cache_signature.json'
    if sigpath.exists():assert json.loads(sigpath.read_text())==sig,'Computation or input changed; use a fresh experiment directory.'
    else:dump(sigpath,sig)
    log('environment',sig)
    if args.stage=='smoke':
        alg=verify_algorithms();X=np.zeros((365,N,25))
        for d in range(1,22):X[d]=features(y,d)
        q,sep=forecast_day(y,21,X)
        # Real-data vertical slice: historical forecasts are past-day persistence.
        pred=np.zeros((365,5,N));pred[1:]=((y[0]-y[1])[:-1])[:,None,:];pred[21]=q
        m=markov(y,pred,21);g,V,label,cands=choose_plan(p,q,m,6000.,float(p.mean()/ETA))
        a=execute((y[0]-y[1])[21],p,g,6000.,m,V)
        dump(OUT/'smoke.json',{'algorithms':alg,'cost':cost(a,p),'chosen':label,'physical':validate(a,(y[0]-y[1])[21],p),'seconds':time.time()-started})
        log('smoke complete',cost(a,p),'seconds',time.time()-started);return
    path=OUT/'forecasts.npz'
    pred=np.load(path)['pred'] if path.exists() else forecasts(y)
    if args.stage=='forecast':log('forecasts complete',time.time()-started);return
    if args.stage=='export':
        runs={s:np.load(OUT/f'{s}.npy') for s in ['main','median_only','greedy_same_plan']}
        k=json.loads((OUT/'controller_selection.json').read_text())['selected']['kappa']
    else:runs,k=backtest(y,p,pred)
    checks=validate_results(y,p,pred,runs,k);data=export(y,p,pred,runs,k,checks)
    preservation=check_preservation()
    manifest={**sig,'runtime':sys.executable,'python':platform.python_version(),'seed':SEED,
        'openpyxl':openpyxl.__version__,'command':[sys.executable,str(Path(__file__).relative_to(ROOT)),'--stage',args.stage],
        'seconds':time.time()-started,'status':'completed, pending user review','original_preservation':preservation,
        'outputs':{str(f.relative_to(BASE)):sha(f) for f in sorted(OUT.glob('*')) if f.is_file() and f.name!='reproduction_manifest.json'}}
    dump(OUT/'reproduction_manifest.json',manifest)
    log('COMPLETE',data['new'],'net forecast',data['forecast'],'preservation',preservation)

if __name__=='__main__':main()
