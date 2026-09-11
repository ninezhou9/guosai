"""第一问：确定性 MILP，输入只读，结果为待审核版。"""
from pathlib import Path
import sys, json, hashlib, csv, platform
from datetime import time
import numpy as np
import scipy
from scipy.optimize import milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix

# 仅复用桌面运行时的纯 Python Excel 包；不修改全局环境。
try:
    import openpyxl
except ModuleNotFoundError:
    sys.path.append(str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/python/Lib/site-packages'))
    import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'代码/results/q1'
FIG = ROOT/'代码/figures/q1'
N = 144
DT = 1/6
ETA = .9
CAP = 5000*DT
TOL = 1e-6

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def clock(m):
    return f'{m//60:02d}:{m%60:02d}'

def load():
    wb = openpyxl.load_workbook(ROOT/'基础数据/附件1.xlsx', read_only=True, data_only=True)
    rows = list(wb.active.values)
    assert len(rows)==145 and len(rows[0])==4
    assert rows[0]==('时间','电价','小区负载','光伏发电预测功率'),rows[0]
    for k,row in enumerate(rows[1:],1):
        v=row[0]
        if isinstance(v,time): minute=v.hour*60+v.minute
        elif v=='0:00+1': minute=1440
        else:
            h,m=map(int,str(v).split(':')); minute=h*60+m
        assert minute==k*10,(k,v)
    a=np.array([r[1:] for r in rows[1:]],float)
    assert a.shape==(144,3) and np.isfinite(a).all() and (a>=0).all()
    assert (a[:,0]>0).all()
    wb.close()
    return a.T

def solve(p,L,V,relax=False):
    # x=[g,c,d,w,E(145),z]；能量变量均为 kWh。
    dim=6*N+1
    G,C,D,W,E,Z=0,N,2*N,3*N,4*N,5*N+1
    cost=np.zeros(dim); cost[G:G+N]=p
    lo=np.zeros(dim); hi=np.full(dim,np.inf)
    hi[C:C+N]=CAP; hi[D:D+N]=CAP
    hi[W:W+N]=V*DT  # 弃光不超过可用光伏
    lo[E:E+N+1]=1200; hi[E:E+N+1]=10800
    hi[Z:Z+N]=1
    a=lil_matrix((4*N+1,dim)); lower=np.full(4*N+1,-np.inf); upper=np.zeros(4*N+1)
    # 第一问为日循环：初末相等，其共同值是决策变量。
    a[4*N,[E,E+N]]=[1,-1]
    lower[4*N]=upper[4*N]=0
    for t in range(N):
        a[t,[G+t,C+t,D+t,W+t]]=[1,-1,1,-1]
        lower[t]=upper[t]=(L[t]-V[t])*DT
        a[N+t,[E+t+1,E+t,C+t,D+t]]=[1,-1,-ETA,1/ETA]
        lower[N+t]=upper[N+t]=0
        a[2*N+t,[C+t,Z+t]]=[1,-CAP]
        a[3*N+t,[D+t,Z+t]]=[1,CAP]; upper[3*N+t]=CAP
    integer=np.zeros(dim); integer[Z:]=0 if relax else 1
    r=milp(cost,integrality=integer,bounds=Bounds(lo,hi),
           constraints=LinearConstraint(a.tocsc(),lower,upper),
           options={'mip_rel_gap':1e-9,'time_limit':120})
    assert r.success,(r.status,r.message)
    return r, [r.x[s:s+N].copy() for s in [G,C,D,W]],r.x[E:E+N+1].copy()

def main():
    OUT.mkdir(parents=True,exist_ok=True); FIG.mkdir(parents=True,exist_ok=True)
    p,L,V=load()
    r,(g,c,d,w),e=solve(p,L,V)
    lp,_,_=solve(p,L,V,relax=True)
    baseline=np.maximum((L-V)*DT,0)
    checks={
        'balance_max_abs_kWh':float(np.max(abs(g+V*DT+d-L*DT-c-w))),
        'state_max_abs_kWh':float(np.max(abs(np.diff(e)-ETA*c+d/ETA))),
        'cyclic_boundary_max_abs_kWh':float(abs(e[0]-e[-1])),
        'soc_bound_violation_kWh':float(max(0,1200-e.min(),e.max()-10800)),
        'power_bound_violation_kWh':float(max(0,c.max()-CAP,d.max()-CAP)),
        'simultaneous_charge_discharge_kWh':float(np.minimum(c,d).max()),
        'nonnegative_violation_kWh':float(max(0,-min(g.min(),c.min(),d.min(),w.min()))),
        'curtailment_bound_violation_kWh':float(max(0,np.max(w-V*DT))),
        'cost_recalculation_error_yuan':float(abs(sum(float(p[i])*float(g[i]) for i in range(N))-r.fun)),
        'MILP_minus_LP_yuan':float(r.fun-lp.fun),
        'daily_energy_error_kWh':float(abs(g.sum()+V.sum()*DT-L.sum()*DT-w.sum()-(1-ETA)*c.sum()-(1/ETA-1)*d.sum()))
    }
    assert all(abs(v)<TOL for k,v in checks.items() if k!='MILP_minus_LP_yuan'),checks
    assert r.fun <= p@baseline+TOL and r.fun >= lp.fun-TOL
    labels=[f'{clock(t*10)}-{clock((t+1)*10)}' for t in range(N)]
    header=['时间段','原始区间结束标签','电价_元每kWh','负载_kW','光伏_kW','购电_kWh','充电_kWh','放电_kWh','弃光_kWh','期初储电_kWh','期末储电_kWh','购电费_元']
    data=[[labels[t],clock((t+1)*10),p[t],L[t],V[t],g[t],c[t],d[t],w[t],e[t],e[t+1],p[t]*g[t]] for t in range(N)]
    with (OUT/'完整调度.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(header);writer.writerows(data)
    wb=openpyxl.Workbook(); ws=wb.active;ws.title='计划购电量'
    ws.append(['时间段','计划购电量（kWh）'])
    for label,val in zip(labels,g):ws.append([label,float(val)])
    ws=wb.create_sheet('充放电量');ws.append(['时间段','充电量（kWh，电网侧）','放电量（kWh，电网侧）'])
    blocks=[]
    for k in range(6):
        row=[f'{clock(k*240)}-{clock((k+1)*240)}',float(c[k*24:(k+1)*24].sum()),float(d[k*24:(k+1)*24].sum())]
        ws.append(row);blocks.append(row)
    ws.append(['0:00储电量（kWh）',float(e[0])]);ws.append(['24:00储电量（kWh）',float(e[-1])])
    ws=wb.create_sheet('表1');ws.append(['时间段','购电量（kWh）'])
    selected=[]
    for hour in [10,12,14,16,18,20]:
        row=[labels[hour*6],float(g[hour*6])];ws.append(row);selected.append(row)
    ws.append(['全天购电量（kWh）',float(g.sum())]);ws.append(['全天购电费（元）',float(p@g)])
    ws=wb.create_sheet('十分钟调度明细');ws.append(header)
    for row in data: ws.append(row)
    ws=wb.create_sheet('口径与状态')
    for row in [
        ['状态','计算审核版；按用户要求恢复结束时标；未冻结'],
        ['时间映射','标签为区间结束时刻；00:10对应00:00-00:10；末行对应23:50-24:00'],
        ['效率','充电、放电效率各0.9；充放电量均为电网侧；往返效率0.81'],
        ['边界','日初与日末相等，共同值由优化决定；全部145状态均在1200-10800 kWh'],
        ['弃光','允许弃光，不售电；第一问不设紧急购电'],
        ['模板','原始官方模板只读；当前结果时段标签已按用户指定结束时标调整，与官方标签有差异']]:ws.append(row)
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    for ws in wb:
        ws.freeze_panes='B2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24536B')
        for column in ws.columns:
            width=min(65,max(18,max(len(str(cell.value or '')) for cell in column)+2))
            ws.column_dimensions[get_column_letter(column[0].column)].width=width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,(int,float)):cell.number_format='0.000000'
                cell.alignment=Alignment(vertical='center',wrap_text=True)
    target=OUT/'result1_审核版.xlsx';wb.save(target)
    reread=openpyxl.load_workbook(target,data_only=True)
    exported=np.array([row[1] for row in list(reread['计划购电量'].values)[1:]])
    assert len(exported)==144 and np.max(abs(exported-g))<TOL
    assert abs(exported@p-r.fun)<TOL
    exported_states=np.array([row[9:11] for row in list(reread['十分钟调度明细'].values)[1:]],dtype=float)
    assert np.max(abs(exported_states[:,0]-e[:-1]))<TOL
    assert np.max(abs(exported_states[:,1]-e[1:]))<TOL
    assert abs(reread['充放电量']['B8'].value-e[0])<TOL
    assert abs(reread['充放电量']['B9'].value-e[-1])<TOL
    for ws in reread:
        assert not any(cell.data_type=='e' for row in ws for cell in row)
    reread.close()
    # 覆盖先前开始时标版本，避免同目录留下两套冲突结果。
    official_path=OUT/'result1.xlsx'
    wb.save(official_path)
    checks['xlsx_roundtrip_max_abs_kWh']=float(np.max(abs(exported-g)))
    plt.rcParams.update({'font.family':'Microsoft YaHei','font.size':10,'axes.unicode_minus':False})
    fig,axes=plt.subplots(3,1,figsize=(7.2,7.4),sharex=True,layout='constrained')
    x=np.arange(N)*DT
    axes[0].step(x,L-V,where='post',label='净负载',color='#3C5488')
    axes[0].step(x,g/DT,where='post',label='购电功率',color='#E64B35')
    axes[0].set_ylabel('功率（kW）');axes[0].legend(ncol=2)
    axes[1].bar(x,c/DT,width=DT,align='edge',label='充电',color='#4DBBD5')
    axes[1].bar(x,-d/DT,width=DT,align='edge',label='放电（负向）',color='#E64B35')
    axes[1].set_ylabel('功率（kW）');axes[1].legend(ncol=2)
    axes[2].plot(np.arange(N+1)*DT,e,color='#3C5488',label='储电量')
    for y in [1200,10800]:axes[2].axhline(y,ls='--',color='gray',lw=.8)
    axes[2].set_ylabel('储电量（kWh）');axes[2].set_xlabel('时刻（h）');axes[2].legend()
    axes[2].set_xlim(0,24);axes[2].set_xticks(range(0,25,4))
    for ax in axes:ax.grid(axis='y',alpha=.2)
    for ext in ['png','svg']:fig.savefig(FIG/f'q1_dispatch.{ext}',dpi=300)
    plt.close(fig)
    summary={'optimal_cost_yuan':float(r.fun),'purchase_kWh':float(g.sum()),'no_storage_cost_yuan':float(p@baseline),
        'saving_yuan':float(p@baseline-r.fun),'saving_percent':float((1-r.fun/(p@baseline))*100),
        'no_storage_purchase_kWh':float(baseline.sum()),'charge_kWh':float(c.sum()),'discharge_kWh':float(d.sum()),
        'curtailment_kWh':float(w.sum()),'load_kWh':float(L.sum()*DT),'pv_kWh':float(V.sum()*DT),
        'initial_soc_kWh':float(e[0]),'terminal_soc_kWh':float(e[-1]),
        'soc_min_kWh':float(e.min()),'soc_max_kWh':float(e.max()),'LP_lower_bound_yuan':float(lp.fun),
        'solver_message':r.message,'mip_gap':float(r.mip_gap),'table1':selected,'table2':blocks,'validation':checks}
    (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    report=['# 第一问计算结果（待用户审核）','',
        f'购电费 **{r.fun:.6f} 元**，全天购电量 **{g.sum():.6f} kWh**。',
        f'不使用储能基线费用 {p@baseline:.6f} 元；节省 {p@baseline-r.fun:.6f} 元（{summary["saving_percent"]:.4f}%）。',
        '基线允许弃光、无售电；所有结果均基于附件1预测值，不代表未知实际发电下的执行保证。','',
        '## 表1 指定时段购电量','|时间段|购电量（kWh）|','|---|---:|']
    report += [f'|{label}|{val:.6f}|' for label,val in selected]
    report += ['','## 表2 分段充放电量（电网侧）','|时间段|充电（kWh）|放电（kWh）|','|---|---:|---:|']
    report += [f'|{label}|{cv:.6f}|{dv:.6f}|' for label,cv,dv in blocks]
    report += ['',f'日初储电量 {e[0]:.6f} kWh，日末储电量 {e[-1]:.6f} kWh。两者仅约束相等，共同值由优化决定，未固定为6000 kWh。','',
        f'LP 松弛下界 {lp.fun:.9f} 元；MILP 与下界差 {r.fun-lp.fun:.3g} 元。这验证当前目标的最优性，不保证调度唯一。',
        '低价充电、高价放电的解释受容量、功率及效率共同约束；不能仅按电价排序独立安排各时段。',
        '','## 数值验收','```json',json.dumps(checks,ensure_ascii=False,indent=2),'```','',
        '所有物理残差按 1e-6 kWh 检查，费用复算按 1e-6 元检查；Excel 写入后重新读取校验。',
        '','## 尚未完成',
        '- 按用户要求恢复结束时标；result1.xlsx与审核版均为此口径，不能声称与官方模板标签一致。',
        '- 0:10映射00:00-00:10，0:00+1映射23:50-24:00；原始附件及官方模板未修改。',
        '- 效率仍各90%；初末仅相等，不固定6000；第二问沿用原结束时标回测。结果待审核、未冻结。']
    (OUT/'第一问结果审核.md').write_text('\n'.join(report),encoding='utf-8')
    command=f'& "{sys.executable}" "代码/q1_solution.py"'
    manifest={'runtime':sys.executable,'python':platform.python_version(),'versions':{m.__name__:m.__version__ for m in [np,scipy,openpyxl,matplotlib]},
        'command_from_project_root':command,'random_seed':None,'deterministic':True,
        'parameters':{'dt_hours':DT,'eta_charge':ETA,'eta_discharge':ETA,'power_kW':5000,'soc_bounds_kWh':[1200,10800],'boundary_condition':'E_0 = E_144; common value optimized','tolerance':TOL,'mip_rel_gap':1e-9},
        'inputs':{str(p.relative_to(ROOT)):sha(p) for p in [ROOT/'基础数据/附件1.xlsx',ROOT/'基础数据/C题.pdf',Path(__file__)]},
        'outputs':{str(p.relative_to(ROOT)):sha(p) for p in [target,official_path,OUT/'summary.json',OUT/'完整调度.csv',OUT/'第一问结果审核.md',FIG/'q1_dispatch.png',FIG/'q1_dispatch.svg']},
        'status':'真实计算完成；按用户指定结束时标映射，存在官方模板标签差异；结果待审核','validation':checks}
    (OUT/'复现清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
